from pathlib import Path
import json
import mimetypes
import os
import platform
import re
import subprocess
import sqlite3
import shutil
import sys
import threading
import traceback
import time

from fastapi import FastAPI, HTTPException, Body, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response
from sqlalchemy import func, or_, Integer
from fastapi.staticfiles import StaticFiles

try:
    import cv2
except ImportError:
    cv2 = None

try:
    import fitz
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

try:
    from backend.app.database import SessionLocal, FileIndex
    from backend.app.config import load_config, save_config
    from backend.app.indexer import start as start_indexing, STATE as STATE
    from backend.app.ai_database import init_ai_database
except ModuleNotFoundError:
    from database import SessionLocal, FileIndex
    from config import load_config, save_config
    from indexer import start as start_indexing, STATE as STATE

app = FastAPI()

APP_SHUTTING_DOWN = False

@app.on_event("shutdown")
def graceful_os_shutdown():
    global APP_SHUTTING_DOWN
    APP_SHUTTING_DOWN = True
    import time
    import logging
    try:
        from backend.app import indexer
    except ImportError:
        try:
            import indexer
        except ImportError:
            return

    try:
        if getattr(indexer, 'logging_enabled', False):
            logging.info("OS Shutdown / App termination detected. Stopping scanners gracefully...")
    except Exception:
        pass

    if hasattr(indexer, 'indexer_stopped'): indexer.indexer_stopped = True
    if hasattr(indexer, 'face_scanner_stopped'): indexer.face_scanner_stopped = True
    if hasattr(indexer, 'object_scanner_stopped'): indexer.object_scanner_stopped = True
    if hasattr(indexer, 'hasher_stopped'): indexer.hasher_stopped = True
    if hasattr(indexer, 'combined_scanner_stopped'): indexer.combined_scanner_stopped = True

    for _ in range(30):
        is_running = any([
            getattr(indexer, 'running', False),
            getattr(indexer, 'combined_scanner_running', False),
            getattr(indexer, 'face_scanner_running', False),
            getattr(indexer, 'object_scanner_running', False),
            getattr(indexer, 'hasher_running', False),
        ])
        if not is_running:
            break
        time.sleep(1)

if sys.platform == 'win32':
    try:
        import win32api
        def console_ctrl_handler(ctrl_type):
            if ctrl_type in (2, 5, 6): 
                graceful_os_shutdown()
                return True
            return False
        win32api.SetConsoleCtrlHandler(console_ctrl_handler, True)
    except ImportError:
        pass

def get_log_path() -> Path:
    if getattr(sys, 'frozen', False):
        return Path(sys.executable).parent / "wabs.log"
    else:
        return Path(__file__).resolve().parent.parent.parent / "wabs.log"

LOGGING_ENABLED = False
LOG_FILE_PATH = get_log_path()

class PrintLogger:
    def __init__(self, stream, max_bytes=5*1024*1024):
        self.stream = stream
        self.max_bytes = max_bytes
        self.current_size = 0
        try:
            if LOG_FILE_PATH.exists():
                self.current_size = LOG_FILE_PATH.stat().st_size
        except Exception:
            pass

    def rotate(self):
        try:
            if LOG_FILE_PATH.exists():
                backup_path = LOG_FILE_PATH.with_suffix('.1.log')
                if backup_path.exists():
                    backup_path.unlink()
                LOG_FILE_PATH.rename(backup_path)
            self.current_size = 0
        except Exception:
            pass

    def write(self, data):
        if not data:
            return
        if self.stream:
            self.stream.write(data)
            self.stream.flush()
        if LOGGING_ENABLED:
            try:
                data_len = len(data.encode('utf-8'))
                if self.current_size + data_len > self.max_bytes:
                    self.rotate()
                with open(LOG_FILE_PATH, "a", encoding="utf-8") as f:
                    f.write(data)
                self.current_size += data_len
            except Exception:
                pass

    def flush(self):
        if self.stream:
            self.stream.flush()

    def isatty(self):
        if self.stream and hasattr(self.stream, 'isatty'):
            return self.stream.isatty()
        return False

sys.stdout = PrintLogger(sys.stdout)
sys.stderr = PrintLogger(sys.stderr)

# Removed heavy face recognition dependencies.
# We will now use cv2 (OpenCV) exclusively for face detection and recognition.

face_scanner_thread = None
face_scanner_running = False

from collections import OrderedDict

class LRUCache:
    def __init__(self, capacity: int):
        self.cache = OrderedDict()
        self.capacity = capacity
        self.lock = threading.Lock()

    def get(self, key):
        with self.lock:
            if key not in self.cache:
                return None
            self.cache.move_to_end(key)
            return self.cache[key]

    def put(self, key, value):
        with self.lock:
            self.cache[key] = value
            self.cache.move_to_end(key)
            if len(self.cache) > self.capacity:
                self.cache.popitem(last=False)
            
    def pop(self, key, default=None):
        with self.lock:
            return self.cache.pop(key, default)

EXEMPLAR_CACHE = LRUCache(capacity=50)

def _evaluate_image_faces(file_path: Path, yunet_path: str):
    if APP_SHUTTING_DOWN:
        return []
    import numpy as np
    try:
        img_array = np.fromfile(str(file_path), np.uint8)
        img = cv2.imdecode(img_array, cv2.IMREAD_COLOR)
        if img is None:
            return []
        
        height, width, _ = img.shape
        target_dim = 800
        scale = 1.0
        if max(height, width) > target_dim:
            scale = target_dim / max(height, width)
            new_w, new_h = int(width * scale), int(height * scale)
            det_img = cv2.resize(img, (new_w, new_h))
        else:
            det_img = img

        detector = cv2.FaceDetectorYN.create(yunet_path, "", (det_img.shape[1], det_img.shape[0]))
        success, faces = detector.detect(det_img)
        
        results = []
        if faces is not None:
            if scale != 1.0:
                faces[:, :14] /= scale
            for face in faces:
                x, y, w, h = [int(v) for v in face[:4]]
                x, y = max(0, x), max(0, y)
                face_area = w * h
                face_crop = img[y:y+h, x:x+w]
                sharpness = 0.0
                if face_crop.size > 0:
                    gray = cv2.cvtColor(face_crop, cv2.COLOR_BGR2GRAY)
                    sharpness = cv2.Laplacian(gray, cv2.CV_64F).var()
                results.append({
                    "area": face_area,
                    "sharpness": sharpness,
                    "score": float(np.sqrt(face_area)) * sharpness if face_area > 0 else 0.0
                })
        return results
    except Exception:
        return []

@app.middleware("http")
async def add_cache_headers(request: Request, call_next):
    response = await call_next(request)
    if request.method == "GET" and response.status_code == 200:
        path = request.url.path
        if path.endswith("/thumbnail") or "/preview/" in path:
            response.headers["Cache-Control"] = "public, max-age=31536000, immutable"
    return response

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"]
)
from pydantic import BaseModel

@app.on_event("startup")
def startup_event():
    global LOGGING_ENABLED
    import logging
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    try:
        cfg = load_config()
        LOGGING_ENABLED = cfg.get("enable_logging", False)
    except Exception:
        pass
    try:
        with SessionLocal() as s:
            count = s.query(FileIndex).count()
            STATE["indexed"] = count
            STATE["current"] = count
    except Exception as e:
        print(f"CRITICAL: Startup database connection failed: {e}")
        traceback.print_exc()

def _parse_json(value):
    if not value:
        return {}
    try:
        return json.loads(value)
    except Exception:
        return {}


def _build_item(r, cache_flag=""):
    v = str(r.modified).replace(" ", "_").replace(":", "") if r.modified else "0"
    return {
        "id": r.id,
        "filename": r.filename,
        "path": r.path,
        "category": r.category,
        "size": r.size,
        "modified": r.modified,
        "extension": r.extension,
        "tags": r.tags,
        "metadata": _parse_json(r.metadata_json),
        "thumbnail": f"/preview/{r.id}?v={v}{cache_flag}"
    }


def get_ai_db_path() -> Path:
    cfg = load_config()
    db_path = cfg.get("database_path")
    if not db_path:
        p = Path("archive.db")
    else:
        p = Path(db_path)
        
    if not p.is_absolute():
        if getattr(sys, 'frozen', False):
            p = Path(sys.executable).parent / p
        else:
            p = Path(__file__).resolve().parent.parent.parent / p
            
    if p.parent.is_file():
        return p.parent.parent / "ai_metadata.db"
    return p.parent / "ai_metadata.db"

def _resolve_path(original_path: Path) -> Path:
    cfg = load_config()
    backup_configs = cfg.get("backup_configs", [])

    for config in backup_configs:
        if config.get("path_mapping_enabled"):
            source_root_str = config.get("backup_path")
            target_root_str = config.get("mapped_backup_path")

            if source_root_str and target_root_str:
                try:
                    # Pre-normalize slashes to ensure Windows network paths (\\) 
                    # parse correctly into .parts even if running on Linux/Mac (/)
                    normalized_orig = str(original_path).replace('\\', os.sep).replace('/', os.sep)
                    normalized_src = source_root_str.replace('\\', os.sep).replace('/', os.sep)
                    
                    src_path = Path(normalized_src)
                    tgt_path = Path(target_root_str)
                    
                    orig_parts = Path(normalized_orig).parts
                    src_parts = src_path.parts
                    
                    if len(orig_parts) >= len(src_parts):
                        # Verify the components actually match before splicing
                        match = True
                        for i in range(len(src_parts)):
                            orig_part = orig_parts[i]
                            src_part = src_parts[i]
                            if platform.system() == "Windows":
                                orig_part = orig_part.lower()
                                src_part = src_part.lower()
                            if orig_part != src_part:
                                match = False
                                break
                        if match:
                            return tgt_path.joinpath(*orig_parts[len(src_parts):])
                except Exception as e:
                    print(f"WARNING: Path remapping failed for {original_path}: {e}")
                    traceback.print_exc()

    # This part is reached if remapping is OFF, or if the path was not applicable for remapping.
    return original_path

def get_bundled_model_path(model_filename: str) -> str:
    if hasattr(sys, '_MEIPASS'):
        # PyInstaller extracts bundled files to a temporary _MEIPASS folder
        return str(Path(sys._MEIPASS) / "backend" / model_filename)
    # Development mode: resolve relative to backend directory
    return str(Path(__file__).parent.parent / model_filename)


def _parse_regex_pattern(query):
    if len(query) < 3 or not query.startswith("/"):
        return None
    last_slash = query.rfind("/")
    if last_slash == 0:
        return None

    pattern = query[1:last_slash]
    flags = query[last_slash + 1:]
    re_flags = 0
    if "i" in flags:
        re_flags |= re.IGNORECASE
    try:
        return re.compile(pattern, re_flags)
    except re.error:
        return None


def _build_search_query(query, s, q_base=None):
    if q_base is None:
        q_base = s.query(FileIndex)
    query = query.strip()
    if not query:
        return q_base

    def text_filter(field, term):
        return func.lower(func.coalesce(field, "")).contains(term)

    raw_tokens = re.findall(r'(?:[^\s"]|"(?:\\.|[^"])*")+', query)
    tokens = [t.replace('"', '') for t in raw_tokens]
    filters = []
    tag_tokens = []
    and_tag_tokens = []
    specific_filters = []
    exclude_filters = []

    for token in tokens:
        lower_token = token.lower()
        if lower_token.startswith("-") and len(lower_token) > 1:
            val = lower_token[1:]
            if val.startswith("tag:"):
                tag_val = token[len("-tag:"):]
                exclude_filters.append(or_(
                    text_filter(FileIndex.tags, tag_val.lower()),
                    text_filter(FileIndex.path, tag_val.lower()),
                    text_filter(FileIndex.filename, tag_val.lower()),
                    text_filter(FileIndex.metadata_json, tag_val.lower())
                ))
            elif val.startswith("object:"):
                exclude_filters.append(or_(
                    text_filter(FileIndex.tags, val),
                    text_filter(FileIndex.path, val),
                    text_filter(FileIndex.filename, val),
                    text_filter(FileIndex.metadata_json, val)
                ))
            elif val.startswith("person:"):
                exclude_filters.append(or_(
                    text_filter(FileIndex.tags, val),
                    text_filter(FileIndex.path, val),
                    text_filter(FileIndex.filename, val),
                    text_filter(FileIndex.metadata_json, val)
                ))
            elif val.startswith("type:"):
                t_val = val[len("type:"):]
                t_val_ext = t_val if t_val.startswith(".") else "." + t_val
                exclude_filters.append(or_(
                    func.lower(FileIndex.extension) == t_val_ext,
                    func.lower(FileIndex.category) == t_val
                ))
            elif val.startswith("name:"):
                n_val = val[len("name:"):]
                exclude_filters.append(text_filter(FileIndex.filename, n_val))
            elif val.startswith("camera:"):
                c_val = val[len("camera:"):]
                exclude_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, '$.camera')).contains(c_val))
            elif val.startswith("resolution:"):
                r_val = val[len("resolution:"):]
                exclude_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, '$.resolution')).contains(r_val))
            elif val.startswith("artist:"):
                a_val = val[len("artist:"):]
                exclude_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, '$.artist')).contains(a_val))
            elif val.startswith("album:"):
                al_val = val[len("album:"):]
                exclude_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, '$.album')).contains(al_val))
            elif val.startswith("genre:"):
                g_val = val[len("genre:"):]
                exclude_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, '$.genre')).contains(g_val))
            elif val.startswith("meta:"):
                orig_val = token[1:]
                parts = orig_val[5:].split(":", 1)
                if len(parts) == 2:
                    exclude_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, f'$.{parts[0]}')).contains(parts[1].lower()))
            else:
                exclude_filters.append(or_(
                    text_filter(FileIndex.filename, val),
                    text_filter(FileIndex.path, val),
                    text_filter(FileIndex.tags, val),
                    text_filter(FileIndex.metadata_json, val)
                ))
            continue
        if lower_token.startswith("date:"):
            val = lower_token[len("date:"):]
            # Date range like YYYY-YYYY
            if m_range := re.match(r"^(\d{4})-(\d{4})$", val):
                start_year, end_year = m_range.groups()
                if int(start_year) <= int(end_year):
                    specific_filters.append(FileIndex.modified >= f"{start_year}-01-01")
                    specific_filters.append(FileIndex.modified < f"{int(end_year) + 1}-01-01")
            # Intelligently parse mm/dd/yyyy or dd/mm/yyyy
            elif m := re.match(r"^(\d{1,2})[-/](\d{1,2})[-/](\d{4})$", val):
                p1, p2, year = m.groups()
                p1, p2 = p1.zfill(2), p2.zfill(2)
                specific_filters.append(or_(
                    text_filter(FileIndex.modified, f"{year}-{p1}-{p2}"),
                    text_filter(FileIndex.modified, f"{year}-{p2}-{p1}")
                ))
            else:
                specific_filters.append(text_filter(FileIndex.modified, val))
        elif lower_token.startswith("+tag:"):
            and_tag_tokens.append(token[len("+tag:"):].lower())
        elif lower_token.startswith("tag:"):
            tag_tokens.append(token[len("tag:"):].lower())
        elif lower_token.startswith("+object:"):
            and_tag_tokens.append(lower_token[1:])
        elif lower_token.startswith("object:"):
            tag_tokens.append(lower_token)
        elif lower_token.startswith("+person:"):
            and_tag_tokens.append(lower_token[1:])
        elif lower_token.startswith("person:"):
            tag_tokens.append(lower_token)
        elif lower_token.startswith("type:"):
            val = lower_token[len("type:"):]
            val_ext = val if val.startswith(".") else "." + val
            specific_filters.append(or_(
                func.lower(FileIndex.extension) == val_ext,
                func.lower(FileIndex.category) == val
            ))
        elif lower_token.startswith("name:"):
            val = lower_token[len("name:"):]
            specific_filters.append(text_filter(FileIndex.filename, val))
        elif lower_token.startswith("size:"):
            val = lower_token[len("size:"):]
            operator = ""
            if val.startswith(">="):
                operator, val = ">=", val[2:]
            elif val.startswith("<="):
                operator, val = "<=", val[2:]
            elif val.startswith(">"):
                operator, val = ">", val[1:]
            elif val.startswith("<"):
                operator, val = "<", val[1:]
                
            bytes_val = None
            if m := re.match(r"^(\d+(?:\.\d+)?)\s*([kmgtp]?b)?$", val.lower()):
                num, unit = float(m.group(1)), m.group(2)
                mult = 1
                if unit == "kb": mult = 1024
                elif unit == "mb": mult = 1024**2
                elif unit == "gb": mult = 1024**3
                elif unit == "tb": mult = 1024**4
                elif unit == "pb": mult = 1024**5
                bytes_val = int(num * mult)
                
            if operator and bytes_val is not None:
                size_col = func.cast(FileIndex.size, Integer)
                if operator == ">=": specific_filters.append(size_col >= bytes_val)
                elif operator == "<=": specific_filters.append(size_col <= bytes_val)
                elif operator == ">": specific_filters.append(size_col > bytes_val)
                elif operator == "<": specific_filters.append(size_col < bytes_val)
            else:
                specific_filters.append(func.lower(FileIndex.size).like(f"{val}%"))
        elif lower_token.startswith("length:"):
            val = lower_token[len("length:"):]
            operator = ""
            if val.startswith(">="):
                operator, val = ">=", val[2:]
            elif val.startswith("<="):
                operator, val = "<=", val[2:]
            elif val.startswith(">"):
                operator, val = ">", val[1:]
            elif val.startswith("<"):
                operator, val = "<", val[1:]
                
            num_val = None
            if m := re.match(r"^(\d+(?:\.\d+)?)\s*([smh])?$", val.lower()):
                num, unit = float(m.group(1)), m.group(2)
                mult = 1
                if unit == "m": mult = 60
                elif unit == "h": mult = 3600
                num_val = num * mult
                
            if operator and num_val is not None:
                duration_col = func.cast(func.json_extract(FileIndex.metadata_json, '$.duration'), Integer)
                fmt_duration_col = func.cast(func.json_extract(FileIndex.metadata_json, '$.format.duration'), Integer)
                length_col = func.cast(func.json_extract(FileIndex.metadata_json, '$.length'), Integer)
                val_col = func.coalesce(duration_col, fmt_duration_col, length_col)
                
                if operator == ">=": specific_filters.append(val_col >= num_val)
                elif operator == "<=": specific_filters.append(val_col <= num_val)
                elif operator == ">": specific_filters.append(val_col > num_val)
                elif operator == "<": specific_filters.append(val_col < num_val)
            else:
                specific_filters.append(text_filter(FileIndex.metadata_json, val))
        elif lower_token.startswith("camera:"):
            val = lower_token[len("camera:"):]
            specific_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, '$.camera')).contains(val))
        elif lower_token.startswith("resolution:"):
            val = lower_token[len("resolution:"):]
            specific_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, '$.resolution')).contains(val))
        elif lower_token.startswith("artist:"):
            val = lower_token[len("artist:"):]
            specific_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, '$.artist')).contains(val))
        elif lower_token.startswith("album:"):
            val = lower_token[len("album:"):]
            specific_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, '$.album')).contains(val))
        elif lower_token.startswith("genre:"):
            val = lower_token[len("genre:"):]
            specific_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, '$.genre')).contains(val))
        elif lower_token.startswith("meta:"):
            parts = token[5:].split(":", 1)
            if len(parts) == 2:
                specific_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, f'$.{parts[0]}')).contains(parts[1].lower()))
        elif lower_token.startswith("fps:"):
            val = lower_token[len("fps:"):]
            operator = ""
            if val.startswith(">="): operator, val = ">=", val[2:]
            elif val.startswith("<="): operator, val = "<=", val[2:]
            elif val.startswith(">"): operator, val = ">", val[1:]
            elif val.startswith("<"): operator, val = "<", val[1:]
            
            try:
                num_val = float(val)
                fps_col = func.json_extract(FileIndex.metadata_json, '$.fps')
                if operator == ">=": specific_filters.append(fps_col >= num_val)
                elif operator == "<=": specific_filters.append(fps_col <= num_val)
                elif operator == ">": specific_filters.append(fps_col > num_val)
                elif operator == "<": specific_filters.append(fps_col < num_val)
                else: specific_filters.append(fps_col == num_val)
            except ValueError:
                specific_filters.append(func.lower(func.json_extract(FileIndex.metadata_json, '$.fps')).contains(val))
        elif "*" in lower_token:
            like_val = lower_token.replace("*", "%")
            specific_filters.append(func.lower(func.coalesce(FileIndex.filename, "")).like(like_val))
        else:
            filters.append(lower_token)

    q = q_base
    for sf in specific_filters:
        q = q.filter(sf)

    if filters:
        for term in filters:
            q = q.filter(or_(
                text_filter(FileIndex.filename, term),
                text_filter(FileIndex.path, term),
                text_filter(FileIndex.tags, term),
                text_filter(FileIndex.metadata_json, term)
            ))
    if tag_tokens:
        q = q.filter(or_(*[
            or_(
                text_filter(FileIndex.tags, tag),
                text_filter(FileIndex.path, tag),
                text_filter(FileIndex.filename, tag),
                text_filter(FileIndex.metadata_json, tag)
            )
            for tag in tag_tokens
        ]))
    if and_tag_tokens:
        for tag in and_tag_tokens:
            q = q.filter(or_(
                text_filter(FileIndex.tags, tag),
                text_filter(FileIndex.path, tag),
                text_filter(FileIndex.filename, tag),
                text_filter(FileIndex.metadata_json, tag)
            ))
    if exclude_filters:
        for ef in exclude_filters:
            q = q.filter(~ef)
    return q

@app.get("/files")
def files(category:str="all", offset:int=0, limit:int=50):
    cfg = load_config()
    ui_prefs = cfg.get("ui_preferences") or {}
    cache_enabled = cfg.get("enable_photo_thumbnail_cache")
    if cache_enabled is None:
        cache_enabled = ui_prefs.get("enable_photo_thumbnail_cache", False)
    cache_flag = "&tc=1" if cache_enabled else ""

    with SessionLocal() as s:
        q = s.query(FileIndex.id, FileIndex.filename, FileIndex.path, FileIndex.category, FileIndex.size, FileIndex.modified, FileIndex.extension, FileIndex.tags, FileIndex.metadata_json)
        if category != "all":
            if category == "other":
                standard = ['photo', 'video', 'audio', 'document', 'ebook', 'code', 'font', 'database', 'compressed', 'installer', 'binary']
                q = q.filter(~FileIndex.category.in_(standard))
            elif category == "duplicates":
                dup_sizes = s.query(FileIndex.size).filter(FileIndex.size != '0', FileIndex.size.isnot(None)).group_by(FileIndex.size).having(func.count(FileIndex.id) > 1)
                q = q.filter(FileIndex.size.in_(dup_sizes))
                q = q.order_by(func.cast(FileIndex.size, Integer).desc(), FileIndex.id)
            else:
                q = q.filter(FileIndex.category == category)
        rows = q.offset(offset).limit(limit).all()
        return [_build_item(r, cache_flag) for r in rows]

@app.get("/search")
def search(query:str="", category:str="all", offset:int=0, limit:int=50):
    cfg = load_config()
    ui_prefs = cfg.get("ui_preferences") or {}
    cache_enabled = cfg.get("enable_photo_thumbnail_cache")
    if cache_enabled is None:
        cache_enabled = ui_prefs.get("enable_photo_thumbnail_cache", False)
    cache_flag = "&tc=1" if str(cache_enabled).lower() in ("true", "1", "yes") else ""

    from sqlalchemy import text
    with SessionLocal() as s:
        q_base = s.query(FileIndex.id, FileIndex.filename, FileIndex.path, FileIndex.category, FileIndex.size, FileIndex.modified, FileIndex.extension, FileIndex.tags, FileIndex.metadata_json)
        if category != "all":
            if category == "other":
                standard = ['photo', 'video', 'audio', 'document', 'ebook', 'code', 'font', 'database', 'compressed', 'installer', 'binary']
                q_base = q_base.filter(~FileIndex.category.in_(standard))
            elif category == "duplicates":
                dup_sizes = s.query(FileIndex.size).filter(FileIndex.size != '0', FileIndex.size.isnot(None)).group_by(FileIndex.size).having(func.count(FileIndex.id) > 1)
                q_base = q_base.filter(FileIndex.size.in_(dup_sizes))
                q_base = q_base.order_by(func.cast(FileIndex.size, Integer).desc(), FileIndex.id)
            else:
                q_base = q_base.filter(FileIndex.category == category)

        query = query.strip()
        if not query:
            rows = q_base.offset(offset).limit(limit).all()
            return [_build_item(r, cache_flag) for r in rows]

        regex = _parse_regex_pattern(query)
        if regex:
            filtered = []
            match_count = 0
            # yield_per prevents loading millions of rows into memory at once
            for r in q_base.yield_per(1000):
                haystack = f"{r.filename or ''} {r.path or ''} {r.tags or ''} {r.metadata_json or ''}"
                if regex.search(haystack):
                    if match_count >= offset:
                        filtered.append(r)
                    match_count += 1
                    if len(filtered) == limit:
                        break
            return [_build_item(r, cache_flag) for r in filtered]

        # Fallback to standard builder for custom search parameters (date:, size:, etc)
        search_prefixes = [
            "date:", "tag:", "type:", "name:", "size:", "length:", "object:", "person:",
            "camera:", "resolution:", "fps:", "artist:", "album:", "genre:", "meta:"
        ]
        if any(prefix in query.lower() for prefix in search_prefixes) or "*" in query or query.startswith("-") or " -" in query or query.startswith("+") or " +" in query:
            q = _build_search_query(query, s, q_base)
            rows = q.offset(offset).limit(limit).all()
            return [_build_item(r, cache_flag) for r in rows]
            
        # FTS5 Lightning-Fast Search Path
        safe_query = query.replace('"', '""').replace("'", "''")
        fts_terms = [f'"{word}" *' for word in safe_query.split() if word]
        if not fts_terms:
            return []
            
        fts_query = " AND ".join(fts_terms)
        
        matching_ids = s.execute(
            text("SELECT rowid FROM files_fts WHERE files_fts MATCH :q ORDER BY rank LIMIT 1000"),
            {"q": fts_query}
        ).scalars().all()

        if not matching_ids:
            return []
            
        rows = q_base.filter(FileIndex.id.in_(matching_ids)).all()
        id_to_row = {r.id: r for r in rows}
        sorted_rows = [id_to_row[i] for i in matching_ids if i in id_to_row]
        
        return [_build_item(r, cache_flag) for r in sorted_rows[offset:offset+limit]]

@app.get("/search/suggestions")
def search_suggestions(q: str = "", limit: int = 5):
    from sqlalchemy import text
    import difflib
    
    q = q.strip().lower()
    if not q:
        return {"type": "none", "suggestions": [], "last_word": ""}
        
    words = q.split()
    last_word = words[-1]
    
    with SessionLocal() as s:
        # 1. Try Auto-complete (fast prefix match) using the FTS5 Vocab table
        results = s.execute(
            text("SELECT term FROM files_fts_vocab WHERE term LIKE :prefix ORDER BY doc DESC LIMIT :limit"),
            {"prefix": f"{last_word}%", "limit": limit}
        ).scalars().all()
        
        if results:
            return {"type": "autocomplete", "suggestions": results, "last_word": last_word}
            
        # 2. Try "Did you mean?" (spell-check fuzzy match) if no exact prefixes exist
        if len(last_word) >= 3:
            all_terms = s.execute(
                text("SELECT term FROM files_fts_vocab WHERE length(term) BETWEEN :min_l AND :max_l ORDER BY doc DESC LIMIT 1000"),
                {"min_l": len(last_word)-2, "max_l": len(last_word)+2}
            ).scalars().all()
            
            close_matches = difflib.get_close_matches(last_word, all_terms, n=limit, cutoff=0.7)
            
            valid_matches = []
            for m in close_matches:
                if m == last_word:
                    continue
                # Filter out FTS Porter Stemming artifacts (e.g., 'purchas' for 'purchase', 'asu' for 'asus')
                if last_word.startswith(m) and len(last_word) - len(m) <= 3:
                    continue
                valid_matches.append(m)
                
            if valid_matches:
                return {"type": "did_you_mean", "suggestions": valid_matches, "last_word": last_word}
                
    return {"type": "none", "suggestions": [], "last_word": last_word}

@app.get("/preview/{item_id}")
def preview(item_id:int, theme: str = "dark"):
    with SessionLocal() as session:
        item = session.get(FileIndex, item_id)
        if not item:
            raise HTTPException(status_code=404, detail="Item not found")

        file_path = _resolve_path(Path(item.path))
        file_category = item.category

    if file_path.exists() and file_path.is_file():
        if file_category == "photo":
            cfg = load_config()
            ui_prefs = cfg.get("ui_preferences") or {}
            cache_enabled = cfg.get("enable_photo_thumbnail_cache")
            if cache_enabled is None:
                cache_enabled = ui_prefs.get("enable_photo_thumbnail_cache", False)
            
            if str(cache_enabled).lower() in ("true", "1", "yes"):
                try:
                    limit_val = cfg.get("photo_thumbnail_size_limit_mb")
                    if limit_val is None:
                        limit_val = ui_prefs.get("photo_thumbnail_size_limit_mb", 5)
                    size_limit_mb = float(limit_val)
                    size_limit_bytes = size_limit_mb * 1024 * 1024
                    
                    if file_path.stat().st_size > size_limit_bytes:
                        thumb_dir = Path(cfg.get("thumbnail_path") or "thumbnails") / ".wabs_cache" / "photos"
                        thumb_dir.mkdir(parents=True, exist_ok=True)
                        cached_thumb = thumb_dir / f"{item_id}.jpg"
                        
                        if cached_thumb.exists():
                            return FileResponse(str(cached_thumb), media_type="image/jpeg")
                            
                        success = False
                        if cv2 is not None:
                            try:
                                import numpy as np
                                img_array = np.fromfile(str(file_path), np.uint8)
                                img = cv2.imdecode(img_array, cv2.IMREAD_COLOR)
                                if img is not None:
                                    height, width = img.shape[:2]
                                    scaling_factor = min(400 / width, 400 / height)
                                    if scaling_factor < 1.0:
                                        new_size = (int(width * scaling_factor), int(height * scaling_factor))
                                        resized_img = cv2.resize(img, new_size, interpolation=cv2.INTER_AREA)
                                    else:
                                        resized_img = img
                                        
                                    is_success, buffer = cv2.imencode(".jpg", resized_img)
                                    if is_success:
                                        with open(str(cached_thumb), "wb") as f:
                                            f.write(buffer.tobytes())
                                        success = True
                            except Exception as e:
                                print(f"OpenCV photo cache failed for {file_path.name}: {e}")
                                
                        if not success:
                            try:
                                from PIL import Image, ImageOps
                                with Image.open(file_path) as pil_img:
                                    pil_img = ImageOps.exif_transpose(pil_img)
                                    if pil_img.mode != 'RGB':
                                        pil_img = pil_img.convert('RGB')
                                    pil_img.thumbnail((400, 400), Image.Resampling.LANCZOS)
                                    pil_img.save(str(cached_thumb), "JPEG", quality=85)
                                    success = True
                            except Exception as e:
                                print(f"Pillow photo cache fallback failed for {file_path.name}: {e}")
                                
                        if success and cached_thumb.exists():
                            return FileResponse(str(cached_thumb), media_type="image/jpeg")
                except Exception as e:
                    print(f"Large photo thumbnail error: {e}")

            media_type, _ = mimetypes.guess_type(str(file_path))
            return FileResponse(str(file_path), media_type=media_type or "application/octet-stream")
        elif file_category == "video":
            cfg = load_config()
            # Enforce an isolated sub-directory so we never overwrite user files
            thumb_dir = Path(cfg.get("thumbnail_path") or "thumbnails") / ".wabs_cache"
            thumb_dir.mkdir(parents=True, exist_ok=True)
            
            cached_thumb = thumb_dir / f"{item_id}.jpg"
            if cached_thumb.exists():
                return FileResponse(str(cached_thumb), media_type="image/jpeg")
            
            if cv2 is not None:
                try:
                    cap = cv2.VideoCapture(str(file_path))
                    frame_count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
                    if frame_count > 0:
                        cap.set(cv2.CAP_PROP_POS_FRAMES, int(frame_count * 0.1)) # Skip to 10% to avoid black start frames
                    success, frame = cap.read()
                    if not success:
                        cap.set(cv2.CAP_PROP_POS_FRAMES, 0)
                        success, frame = cap.read()
                    if success:
                        height, width = frame.shape[:2]
                        scaling_factor = min(400 / width, 300 / height)
                        new_size = (int(width * scaling_factor), int(height * scaling_factor))
                        resized_frame = cv2.resize(frame, new_size, interpolation=cv2.INTER_AREA)
                        is_success, buffer = cv2.imencode(".jpg", resized_frame)
                        if is_success:
                            with open(str(cached_thumb), "wb") as f:
                                f.write(buffer.tobytes())
                        cap.release()
                        if cached_thumb.exists():
                            return FileResponse(str(cached_thumb), media_type="image/jpeg")
                    cap.release()
                except Exception as e:
                    print(f"ERROR: Video thumbnail error for {file_path}: {e}")
                    traceback.print_exc()
                    
        elif file_category in ["document", "code"] or file_path.suffix.lower() in [".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml", ".log"]:
            if file_path.suffix.lower() == ".pdf":
                cfg = load_config()
                # Enforce an isolated sub-directory so we never overwrite user files
                thumb_dir = Path(cfg.get("thumbnail_path") or "thumbnails") / ".wabs_cache"
                thumb_dir.mkdir(parents=True, exist_ok=True)
                
                cached_thumb = thumb_dir / f"{item_id}.jpg"
                if cached_thumb.exists():
                    return FileResponse(str(cached_thumb), media_type="image/jpeg")
                
                if fitz is not None:
                    try:
                        doc = fitz.open(str(file_path))
                        
                        if doc.needs_pass:
                            doc.close()
                            bg_fill = '#f8fafc' if theme == 'light' else '#111827'
                            text_fill_1 = '#0f172a' if theme == 'light' else '#94a3b8'
                            text_fill_2 = '#334155' if theme == 'light' else '#64748b'
                            placeholder = f"""
<svg xmlns='http://www.w3.org/2000/svg' width='400' height='300' viewBox='0 0 400 300'>
  <rect width='400' height='300' fill='{bg_fill}'/>
  <text x='50%' y='45%' fill='{text_fill_1}' font-family='Segoe UI,Arial' font-size='22' text-anchor='middle'>Preview unavailable</text>
  <text x='50%' y='60%' fill='{text_fill_2}' font-family='Segoe UI,Arial' font-size='16' text-anchor='middle'>ENCRYPTED PDF</text>
</svg>
"""
                            return Response(content=placeholder, media_type='image/svg+xml')

                        page = doc.load_page(0)
                        pix = page.get_pixmap(matrix=fitz.Matrix(0.5, 0.5))
                        pix.save(str(cached_thumb))
                        doc.close()
                        if cached_thumb.exists():
                            return FileResponse(str(cached_thumb), media_type="image/jpeg")
                    except Exception as e:
                        print(f"ERROR: PDF thumbnail error for {file_path}: {e}")
                        traceback.print_exc()
            elif file_path.suffix.lower() == ".docx" and docx is not None:
                try:
                    doc = docx.Document(str(file_path))
                    lines = []
                    for p in doc.paragraphs:
                        if p.text.strip():
                            lines.append(p.text.strip())
                        if len(lines) >= 11:
                            break
                    
                    text_fill = '#0f172a' if theme == 'light' else '#cbd5e1'
                    svg_lines = ""
                    y = 28
                    for line in lines:
                        safe_line = line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')[:50]
                        svg_lines += f"<text x='16' y='{y}' fill='{text_fill}' font-family='monospace' font-size='13'>{safe_line}</text>\n"
                        y += 24
                        
                    bg_fill = '#f8fafc' if theme == 'light' else '#0f172a'
                    text_svg = f"<svg xmlns='http://www.w3.org/2000/svg' width='400' height='300' viewBox='0 0 400 300'>\n  <rect width='400' height='300' fill='{bg_fill}'/>\n{svg_lines}</svg>"
                    return Response(content=text_svg, media_type='image/svg+xml')
                except Exception as e:
                    print(f"ERROR: DOCX thumbnail error for {file_path}: {e}")
            else:
                text_extensions = [".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml", ".py", ".js", ".html", ".css", ".c", ".cpp", ".h", ".java", ".cs", ".go", ".rs", ".rb", ".php", ".sh", ".bat", ".sql"]
                if file_path.suffix.lower() in text_extensions:
                    try:
                        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                            lines = [f.readline().rstrip('\n') for _ in range(11)]
                        
                        text_fill = '#0f172a' if theme == 'light' else '#cbd5e1'
                        svg_lines = ""
                        y = 28
                        for line in lines:
                            safe_line = line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')[:50]
                            svg_lines += f"<text x='16' y='{y}' fill='{text_fill}' font-family='monospace' font-size='13'>{safe_line}</text>\n"
                            y += 24
                            
                        bg_fill = '#f8fafc' if theme == 'light' else '#0f172a'
                        text_svg = f"<svg xmlns='http://www.w3.org/2000/svg' width='400' height='300' viewBox='0 0 400 300'>\n  <rect width='400' height='300' fill='{bg_fill}'/>\n{svg_lines}</svg>"
                        return Response(content=text_svg, media_type='image/svg+xml')
                    except Exception as e:
                        print(f"ERROR: Text thumbnail error for {file_path}: {e}")
                        traceback.print_exc()

    bg_fill = '#f8fafc' if theme == 'light' else '#111827'
    text_fill_1 = '#0f172a' if theme == 'light' else '#94a3b8'
    text_fill_2 = '#334155' if theme == 'light' else '#64748b'
    placeholder = f"""
<svg xmlns='http://www.w3.org/2000/svg' width='400' height='300' viewBox='0 0 400 300'>
  <rect width='400' height='300' fill='{bg_fill}'/>
  <text x='50%' y='45%' fill='{text_fill_1}' font-family='Segoe UI,Arial' font-size='22' text-anchor='middle'>Preview unavailable</text>
  <text x='50%' y='60%' fill='{text_fill_2}' font-family='Segoe UI,Arial' font-size='16' text-anchor='middle'>{file_category.upper()}</text>
</svg>
"""
    return Response(content=placeholder, media_type='image/svg+xml')

@app.post("/open/{item_id}")
def open_file(item_id:int):
    with SessionLocal() as session:
        item = session.get(FileIndex, item_id)
        if not item:
            raise HTTPException(status_code=404, detail="Item not found")

        return open_file_path(Path(item.path))

@app.post("/open-path")
def open_file_path(path: str = Body(..., embed=True)):
    file_path = _resolve_path(Path(path))
    if not file_path.exists() or not file_path.is_file():
        raise HTTPException(status_code=404, detail="File not found on disk")

    system_name = platform.system()
    try:
        if system_name == "Windows":
            norm_path = os.path.normpath(file_path)
            subprocess.Popen(f'start "" "{norm_path}"', shell=True)
        elif system_name == "Darwin":
            subprocess.Popen(["open", str(file_path)])
        else:
            subprocess.Popen(["xdg-open", str(file_path)])
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Unable to open file: {exc}")

    return {"opened": True, "path": str(file_path), "platform": system_name}

@app.post("/open-folder")
def open_folder(path: str = Body(..., embed=True)):
    resolved_path = _resolve_path(Path(path))
    folder_path = resolved_path.parent
    if not folder_path.exists() or not folder_path.is_dir():
        raise HTTPException(status_code=404, detail="Folder not found")

    system_name = platform.system()
    try:
        if system_name == "Windows":
            if resolved_path.exists() and resolved_path.is_file():
                norm_path = os.path.normpath(resolved_path)
                subprocess.Popen(['explorer', '/select,', norm_path])
            else:
                norm_folder = os.path.normpath(folder_path)
                subprocess.Popen(['explorer', norm_folder])
        elif system_name == "Darwin":
            if resolved_path.exists() and resolved_path.is_file():
                subprocess.Popen(["open", "-R", str(resolved_path)])
            else:
                subprocess.Popen(["open", str(folder_path)])
        else:
            subprocess.Popen(["xdg-open", str(folder_path)])
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Unable to open folder: {exc}")

    return {"opened": True, "path": str(folder_path), "platform": system_name}

@app.post("/delete-files")
def delete_files(paths: list[str] = Body(..., embed=True)):
    cfg = load_config()
    if cfg.get("read_only_mode", True):
        raise HTTPException(status_code=403, detail="Read-Only Mode is enabled. Deletion is blocked.")

    backup_configs = cfg.get("backup_configs", [])
    for path_str in paths:
        for config in backup_configs:
            bp = config.get("backup_path")
            if bp and config.get("read_only_mode", True):
                orig_norm = os.path.normpath(path_str)
                bp_norm = os.path.normpath(bp)
                if platform.system() == "Windows":
                    if orig_norm.lower().startswith(bp_norm.lower()):
                        raise HTTPException(status_code=403, detail=f"Read-Only Mode is enabled for backup '{config.get('name', 'location')}'. Deletion is blocked.")
                else:
                    if orig_norm.startswith(bp_norm):
                        raise HTTPException(status_code=403, detail=f"Read-Only Mode is enabled for backup '{config.get('name', 'location')}'. Deletion is blocked.")

    deleted_count = 0
    with SessionLocal() as session:
        for path_str in paths:
            file_path = _resolve_path(Path(path_str))
            try:
                if file_path.exists() and file_path.is_file():
                    file_path.unlink() # Deletes file from disk
                session.query(FileIndex).filter(FileIndex.path == path_str).delete()
                deleted_count += 1
            except Exception as e:
                if cfg.get("enable_logging"):
                    import logging
                    logging.error(f"Critical error: Failed to delete {path_str}: {e}", exc_info=True)
                print(f"Failed to delete {path_str}: {e}")
        session.commit()
    if cfg.get("enable_logging"):
        import logging
        logging.info(f"Deleted {deleted_count} files.")
    return {"deleted": deleted_count}

@app.post("/copy-files")
def copy_files(paths: list[str] = Body(...), destination: str = Body(...)):
    dest_path = Path(destination)
    if not dest_path.exists() or not dest_path.is_dir():
        raise HTTPException(status_code=400, detail="Invalid destination directory")
    
    copied_count = 0
    for path_str in paths:
        src = _resolve_path(Path(path_str))
        if src.exists() and src.is_file():
            try:
                shutil.copy2(src, dest_path / src.name)
                copied_count += 1
            except Exception as e:
                if load_config().get("enable_logging"):
                    import logging
                    logging.error(f"Critical error: Failed to copy {path_str}: {e}", exc_info=True)
                print(f"Failed to copy {path_str}: {e}")
    if load_config().get("enable_logging"):
        import logging
        logging.info(f"Successfully copied {copied_count} files to destination.")
    return {"copied": copied_count}

@app.post("/move-files")
def move_files(paths: list[str] = Body(...), destination: str = Body(...)):
    cfg = load_config()
    if cfg.get("read_only_mode", True):
        raise HTTPException(status_code=403, detail="Read-Only Mode is enabled. Moving files is blocked.")

    backup_configs = cfg.get("backup_configs", [])
    for path_str in paths:
        for config in backup_configs:
            bp = config.get("backup_path")
            if bp and config.get("read_only_mode", True):
                orig_norm = os.path.normpath(path_str)
                bp_norm = os.path.normpath(bp)
                if platform.system() == "Windows":
                    if orig_norm.lower().startswith(bp_norm.lower()):
                        raise HTTPException(status_code=403, detail=f"Read-Only Mode is enabled for backup '{config.get('name', 'location')}'. Moving files is blocked.")
                else:
                    if orig_norm.startswith(bp_norm):
                        raise HTTPException(status_code=403, detail=f"Read-Only Mode is enabled for backup '{config.get('name', 'location')}'. Moving files is blocked.")

    dest_path = Path(destination)
    if not dest_path.exists() or not dest_path.is_dir():
        raise HTTPException(status_code=400, detail="Invalid destination directory")
    
    moved_count = 0
    updates = {}
    with SessionLocal() as session:
        for path_str in paths:
            src = _resolve_path(Path(path_str))
            if src.exists() and src.is_file():
                try:
                    new_target = dest_path / src.name
                    shutil.move(str(src), str(new_target))
                    
                    db_item = session.query(FileIndex).filter(FileIndex.path == path_str).first()
                    if db_item:
                        db_item.path = str(new_target)
                    
                    updates[path_str] = str(new_target)
                    moved_count += 1
                except Exception as e:
                    if cfg.get("enable_logging"):
                        import logging
                        logging.error(f"Critical error: Failed to move {path_str}: {e}", exc_info=True)
                    print(f"Failed to move {path_str}: {e}")
        session.commit()
    if cfg.get("enable_logging"):
        import logging
        logging.info(f"Successfully moved {moved_count} files to destination.")
    return {"moved": moved_count, "updates": updates}

@app.get("/stats")
def stats():
    cfg = load_config()
    with SessionLocal() as s:
        # Grouping drastically speeds up counting for millions of rows compared to 5 separate counts
        results = s.query(FileIndex.category, func.count(FileIndex.id)).group_by(FileIndex.category).all()
        stats_dict = {"total": 0, "duplicates": 0, "photos": 0, "videos": 0, "audio": 0, "documents": 0, "ebooks": 0, "code": 0, "fonts": 0, "databases": 0, "compressed": 0, "installers": 0, "binaries": 0, "others": 0}
        for cat, count in results:
            stats_dict["total"] += count
            if cat == "photo": stats_dict["photos"] += count
            elif cat == "video": stats_dict["videos"] += count
            elif cat == "audio": stats_dict["audio"] += count
            elif cat == "document": stats_dict["documents"] += count
            elif cat == "ebook": stats_dict["ebooks"] += count
            elif cat == "code": stats_dict["code"] += count
            elif cat == "font": stats_dict["fonts"] += count
            elif cat == "database": stats_dict["databases"] += count
            elif cat == "compressed": stats_dict["compressed"] += count
            elif cat == "installer": stats_dict["installers"] += count
            elif cat == "binary": stats_dict["binaries"] += count
            else: stats_dict["others"] += count
            
        dup_subq = s.query(func.count(FileIndex.id).label('c')).filter(FileIndex.size != '0', FileIndex.size.isnot(None)).group_by(FileIndex.size).having(func.count(FileIndex.id) > 1).subquery()
        dup_count = s.query(func.sum(dup_subq.c.c)).scalar() or 0
        stats_dict["duplicates"] = int(dup_count)
        
        stats_dict["known_faces"] = 0
        stats_dict["unknown_faces"] = 0
        ai_db_path = get_ai_db_path()
        if ai_db_path.exists():
            try:
                with sqlite3.connect(ai_db_path, timeout=15) as conn:
                    conn.execute("PRAGMA journal_mode=WAL;")
                    cursor = conn.cursor()
                    cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='faces'")
                    if cursor.fetchone():
                        hidden_people = cfg.get("hidden_people")
                        if hidden_people is None:
                            hidden_people = cfg.get("ui_preferences", {}).get("hidden_people", [])
                        if not isinstance(hidden_people, list):
                            hidden_people = []
                        hidden_ids = [str(pid) for pid in hidden_people if str(pid).isdigit()]
                        hidden_clause = f" AND people.id NOT IN ({','.join(hidden_ids)})" if hidden_ids else ""
                        
                        cursor.execute(f"SELECT COUNT(DISTINCT people.id) FROM faces JOIN people ON faces.person_id = people.id WHERE people.name NOT LIKE 'Unknown Person%' {hidden_clause}")
                        stats_dict["known_faces"] = cursor.fetchone()[0] or 0
                        
                        cursor.execute(f"SELECT COUNT(DISTINCT people.id) FROM faces JOIN people ON faces.person_id = people.id WHERE people.name LIKE 'Unknown Person%' {hidden_clause}")
                        stats_dict["unknown_faces"] = cursor.fetchone()[0] or 0
            except Exception as e:
                print(f"Error fetching AI stats: {e}")
                
        return stats_dict

@app.get("/timeline")
def timeline(category: str = "all"):
    with SessionLocal() as s:
        q = s.query(func.date(FileIndex.modified).label("date"), func.count(FileIndex.id))
        if category != "all":
            if category == "other":
                standard = ['photo', 'video', 'audio', 'document', 'ebook', 'code', 'font', 'database', 'compressed', 'installer', 'binary']
                q = q.filter(~FileIndex.category.in_(standard))
            elif category == "duplicates":
                dup_sizes = s.query(FileIndex.size).filter(FileIndex.size != '0', FileIndex.size.isnot(None)).group_by(FileIndex.size).having(func.count(FileIndex.id) > 1)
                q = q.filter(FileIndex.size.in_(dup_sizes))
            else:
                q = q.filter(FileIndex.category == category)
        
        q = q.filter(FileIndex.modified.isnot(None))
        q = q.group_by('date').order_by('date')
        rows = q.all()
        return [{"date": r[0], "count": r[1]} for r in rows if r[0]]

class IndexRequest(BaseModel):
    tag: bool = False
    face: bool = False

@app.post("/indexer/set-options")
def indexer_set_options(req: IndexRequest):
    try:
        cfg = load_config()
        cfg["run_face_scan"] = req.face
        cfg["run_object_scan"] = req.tag
        save_config(cfg)
        return {"saved": True}
    except Exception as e:
        print(f"Warning: could not save scan options to config: {e}")
        raise HTTPException(status_code=500, detail="Could not save options")

@app.get("/indexer/status")
def indexer_status():
    status = dict(STATE)
    status["face_scanner_running"] = face_scanner_running
    status["object_scanner_running"] = object_scanner_running
    status["combined_scanner_running"] = combined_scanner_running
    status["combined_scanner_stopped"] = combined_scanner_stopped
    return status

@app.post("/indexer/start")
def indexer_start(req: IndexRequest = None):
    global combined_scanner_thread, combined_scanner_running, combined_scanner_stopped
    if req is None:
        req = IndexRequest()

    if cv2 is None and (req.tag or req.face):
        raise HTTPException(status_code=500, detail="OpenCV is required for face and object recognition.")
    if STATE.get("running") or combined_scanner_running:
        return {"started": True, "ignored": True}
    STATE["update_only"] = False

    if req.tag or req.face:
        combined_scanner_running = True
        combined_scanner_stopped = False
        combined_scanner_thread = threading.Thread(target=_process_unified_scanners, kwargs={"run_index": True, "run_object": req.tag, "run_face": req.face})
        combined_scanner_thread.start()
    else:
        start_indexing()
    if load_config().get("enable_logging"):
        import logging
        logging.info("Archive indexing started.")
    return {"started": STATE.get("running", True)}

@app.post("/indexer/pause")
def indexer_pause():
    if (STATE.get("running") or combined_scanner_running) and not STATE.get("paused"):
        STATE["paused"] = True
        STATE["status"] = "Paused"
    return STATE

@app.post("/indexer/resume")
def indexer_resume():
    if not STATE.get("running") and not combined_scanner_running:
        # App was closed or stopped. Resume intelligently continues from the DB state.
        STATE["update_only"] = True
        start_indexing()
        return {"resumed_from_db": True}
    if (STATE.get("running") or combined_scanner_running) and STATE.get("paused"):
        STATE["paused"] = False
        STATE["status"] = "Indexing & Scanning..." if combined_scanner_running else "Indexing"
    return STATE

@app.post("/indexer/stop")
def indexer_stop():
    if STATE.get("running") or combined_scanner_running:
        STATE["stopped"] = True
        STATE["paused"] = False
        STATE["status"] = "Stopping..."

    global combined_scanner_stopped, face_scanner_running, object_scanner_running
    combined_scanner_stopped = True
    face_scanner_running = False
    object_scanner_running = False
    if load_config().get("enable_logging"):
        import logging
        logging.info("Archive indexing stopped.")
    return STATE

@app.post("/indexer/update")
def indexer_update(req: IndexRequest = None):
    global combined_scanner_thread, combined_scanner_running, combined_scanner_stopped
    if req is None:
        req = IndexRequest()

    if cv2 is None and (req.tag or req.face):
        raise HTTPException(status_code=500, detail="OpenCV is required for face and object recognition.")
    if STATE.get("running") or combined_scanner_running:
        return {"updating": False, "ignored": True}
    STATE["update_only"] = True

    if req.tag or req.face:
        combined_scanner_running = True
        combined_scanner_stopped = False
        combined_scanner_thread = threading.Thread(target=_process_unified_scanners, kwargs={"run_index": True, "run_object": req.tag, "run_face": req.face})
        combined_scanner_thread.start()
    else:
        start_indexing()
    if load_config().get("enable_logging"):
        import logging
        logging.info("Archive indexing update started.")
    return {"updating": True}

@app.post("/indexer/reindex")
def indexer_reindex(req: IndexRequest = None):
    global combined_scanner_thread, combined_scanner_running, combined_scanner_stopped
    if req is None:
        req = IndexRequest()

    if cv2 is None and (req.tag or req.face):
        raise HTTPException(status_code=500, detail="OpenCV is required for face and object recognition.")
    if STATE.get("running") or combined_scanner_running:
        return {"reindexing": False, "ignored": True}
    with SessionLocal() as s:
        from sqlalchemy import text
        s.query(FileIndex).delete()
        try:
            s.execute(text("DELETE FROM sqlite_sequence WHERE name='files'"))
        except Exception:
            pass
        s.commit()

    cfg = load_config()
    # Safely rmtree ONLY our isolated cache directory, ignoring the parent folder entirely
    thumb_dir = Path(cfg.get("thumbnail_path") or "thumbnails") / ".wabs_cache"
    if thumb_dir.exists() and thumb_dir.is_dir():
        try:
            shutil.rmtree(thumb_dir)
        except Exception as e:
            print(f"Failed to clear thumbnails directory: {e}")

    ai_db_path = get_ai_db_path()
    if ai_db_path.exists():
        try:
            with sqlite3.connect(ai_db_path, timeout=15) as conn:
                conn.execute("PRAGMA journal_mode=WAL;")
                for table in ['faces', 'people', 'processed_files', 'processed_objects']:
                    try:
                        conn.execute(f"DELETE FROM {table}")
                    except sqlite3.OperationalError as e:
                        if "no such table" not in str(e).lower():
                            raise
                try:
                    conn.execute("DELETE FROM sqlite_sequence WHERE name IN ('faces', 'people', 'processed_files', 'processed_objects')")
                except Exception:
                    pass
                conn.commit()
        except Exception as e:
            print(f"Failed to clear AI database: {e}")

    STATE["indexed"] = 0
    STATE["current"] = 0
    STATE["total"] = 0
    STATE["update_only"] = False

    if req.tag or req.face:
        combined_scanner_running = True
        combined_scanner_stopped = False
        combined_scanner_thread = threading.Thread(target=_process_unified_scanners, kwargs={"run_index": True, "run_object": req.tag, "run_face": req.face})
        combined_scanner_thread.start()
    else:
        start_indexing()
    if load_config().get("enable_logging"):
        import logging
        logging.info("Archive re-indexing started.")
    return {"reindexing": True}

@app.get("/choose-path")
def choose_path_api(mode: str = "directory"):
    try:
        import tkinter as tk
        from tkinter import filedialog
        
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True) # Ensures the dialog appears above the browser
        if mode == "directory":
            path = filedialog.askdirectory(parent=root, title="Select Directory")
        else:
            path = filedialog.askopenfilename(parent=root, title="Select File")
        root.destroy()
        return {"path": path or ""}
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Dialog failed: {exc}")

@app.get("/settings")
def settings():
    cfg = load_config()
    if not isinstance(cfg, dict):
        cfg = {}
        
    # Define a schema of all possible settings and their defaults to ensure a complete config
    defaults = {
        "database_path": "archive.db",
        "thumbnail_path": "thumbnails",
        "enable_logging": False,
        "theme": "dark",
        "enable_photo_thumbnail_cache": False,
        "photo_thumbnail_size_limit_mb": 5,
        "allow_unverified_deletion": False,
        "animations_enabled": True,
        "show_full_timeline": False,
        "read_only_mode": True,
        "ai_enabled": False,
        "ai_provider": "",
        "ai_model": "",
        "openai_api_key": "",
        "face_sensitivity": "medium",
        "face_clustering_sensitivity": "medium",
        "object_sensitivity": "medium",
        "min_unknown_photos": 1,
        "run_face_scan": False,
        "run_object_scan": False,
        "backup_configs": [],
        "smart_searches": [],
    }

    # Recursively merge defaults into the loaded config to prevent crashes from missing keys
    def merge_defaults(config, defaults_dict):
        for key, default_value in defaults_dict.items():
            if key not in config:
                config[key] = default_value
            elif isinstance(default_value, dict) and isinstance(config[key], dict):
                merge_defaults(config[key], default_value)
        return config
        
    return merge_defaults(cfg, defaults)

@app.post("/settings")
def save(data:dict):
    global LOGGING_ENABLED
    save_config(data)
    LOGGING_ENABLED = data.get("enable_logging", False)
    # Log settings update
    if load_config().get("enable_logging"):
        import logging
        logging.info("Configuration file updated.")
    return {"saved":True}

@app.post("/clear-cache")
def clear_cache():
    import logging
    cfg = load_config()
    thumb_dir = Path(cfg.get("thumbnail_path") or "thumbnails") / ".wabs_cache"
    if thumb_dir.exists() and thumb_dir.is_dir():
        try:
            shutil.rmtree(thumb_dir)
            if cfg.get("enable_logging"):
                logging.info("Cleared thumbnail cache.")
            return {"cleared": True}
        except Exception as e:
            if cfg.get("enable_logging"):
                logging.error(f"Critical error: Failed to clear cache: {e}", exc_info=True)
            raise HTTPException(status_code=500, detail=f"Failed to clear cache: {e}")
    return {"cleared": True, "message": "Cache was already empty"}

@app.post("/shutdown")
def shutdown(request: Request):
    global APP_SHUTTING_DOWN
    APP_SHUTTING_DOWN = True
    import logging
    logging_enabled = load_config().get("enable_logging")
    # To achieve a graceful shutdown, we access the server instance stored in the app state
    # by run.py and set its `should_exit` flag. This is the production method.
    if hasattr(request.app.state, 'server'):
        server = request.app.state.server
        # We run this in a thread to allow the HTTP response to be sent to the client first.
        def graceful_shutdown():
            time.sleep(2.0) # A delay to ensure the response is sent and files are closed.
            if logging_enabled:
                logging.info("Server is shutting down (Production method).")
            server.should_exit = True
        threading.Thread(target=graceful_shutdown).start()
        return {"shutdown": True, "message": "Server is shutting down..."}
    # Fallback for development mode (e.g., running with `uvicorn main:app --reload`)
    else:
        import os
        import signal
        def dev_shutdown():
            time.sleep(2.0) # A delay to ensure the response is sent and files are closed.
            # This sends a termination signal that Uvicorn's CLI runner will catch
            # and use to perform a graceful shutdown.
            if logging_enabled:
                logging.info("Server shutdown signal sent (Development method).")
            os.kill(os.getpid(), signal.SIGTERM)
        threading.Thread(target=dev_shutdown).start()
        return {"shutdown": True, "message": "Server shutdown signal sent..."}

@app.post("/verify-duplicates")
def verify_duplicates():
    try:
        from backend.app.indexer import background_lazy_hasher
    except ModuleNotFoundError:
        from indexer import background_lazy_hasher
    import threading
    threading.Thread(target=background_lazy_hasher, daemon=True).start()
    return {"status": "started"}

@app.post("/stop-verify-duplicates")
def stop_verify_duplicates():
    try:
        from backend.app.state import STATE
    except ModuleNotFoundError:
        from state import STATE
    STATE["hasher_stopped"] = True
    return {"status": "stopping"}

def _cosine_similarity(vec1, vec2):
    import numpy as np
    v1, v2 = np.array(vec1), np.array(vec2)
    norm_1, norm_2 = np.linalg.norm(v1), np.linalg.norm(v2)
    if norm_1 == 0 or norm_2 == 0:
        return 0.0
    return float(np.dot(v1, v2) / (norm_1 * norm_2))

def _process_unified_scanners(run_index: bool = False, run_face: bool = False, run_object: bool = False):
    global face_scanner_running, object_scanner_running, combined_scanner_running
    global combined_scanner_stopped
    
    if run_face:
        face_scanner_running = True
    if run_object:
        object_scanner_running = True
        STATE["object_scanner_stopped"] = False

    try:
        import numpy as np
        import os
        import time
        try:
            from backend.app.indexer import classify, extract_metadata_for_file, build_tags
        except ModuleNotFoundError:
            from indexer import classify, extract_metadata_for_file, build_tags
            
        cfg = load_config()
        ai_db_path = get_ai_db_path()

        # --- Object Setup ---
        net, classes, object_threshold = None, None, 0.15
        if run_object:
            model_path = get_bundled_model_path("mobilenetv2-small.onnx")
            classes_path = get_bundled_model_path("imagenet_classes.txt")
            if Path(model_path).exists() and Path(classes_path).exists():
                net = cv2.dnn.readNetFromONNX(model_path)
                with open(classes_path, 'rt') as f:
                    classes = [line.strip() for line in f.readlines()]
            object_sensitivity = cfg.get("object_sensitivity", "medium")
            object_threshold = 0.10 if object_sensitivity == "high" else 0.30 if object_sensitivity == "low" else 0.15

        # --- Face Setup ---
        detector, recognizer, clusters, p_count = None, None, {}, 0
        cluster_matrix_norm = None
        cluster_ids_list = []
        face_threshold, cluster_threshold = 0.70, 0.55
        if run_face:
            yunet_path = get_bundled_model_path("face_detection_yunet_2023mar.onnx")
            sface_path = get_bundled_model_path("face_recognition_sface_2021dec.onnx")

            if Path(yunet_path).exists() and Path(sface_path).exists():
                detector = cv2.FaceDetectorYN.create(yunet_path, "", (320, 320))
                face_sensitivity = cfg.get("face_sensitivity", "medium")
                face_threshold = 0.55 if face_sensitivity == "high" else 0.85 if face_sensitivity == "low" else 0.70
                detector.setScoreThreshold(face_threshold)
                recognizer = cv2.FaceRecognizerSF.create(sface_path, "")
            else:
                print("Face recognition models not found. Ensure .onnx files are in the backend folder.")
                return

            cluster_sensitivity = cfg.get("face_clustering_sensitivity", "medium")
            cluster_threshold = 0.65 if cluster_sensitivity == "high" else 0.45 if cluster_sensitivity == "low" else 0.55

        # --- DB Initialization ---
        face_processed_ids = set()
        object_processed_ids = set()

        with sqlite3.connect(ai_db_path, timeout=15) as conn:
            conn.execute("PRAGMA journal_mode=WAL;")
            cursor = conn.cursor()
            if run_face:
                cursor.execute('''CREATE TABLE IF NOT EXISTS people (
                                id INTEGER PRIMARY KEY AUTOINCREMENT,
                                name TEXT DEFAULT 'Unknown Person',
                                thumbnail_file_id INTEGER
                              )''')
                cursor.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_people_name ON people(name)")
                cursor.execute('''CREATE TABLE IF NOT EXISTS faces (
                                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                                    person_id INTEGER,
                                    file_id INTEGER,
                                    embedding_json TEXT,
                                    FOREIGN KEY(person_id) REFERENCES people(id)
                                )''')
                cursor.execute("CREATE INDEX IF NOT EXISTS idx_faces_person_file ON faces(person_id, file_id)")
                cursor.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_faces_unique ON faces(person_id, file_id, embedding_json)")
                cursor.execute('''CREATE TABLE IF NOT EXISTS processed_files (
                                    file_id INTEGER PRIMARY KEY
                                )''')
                cursor.execute("INSERT OR IGNORE INTO processed_files (file_id) SELECT DISTINCT file_id FROM faces")
                cursor.execute("SELECT file_id FROM processed_files")
                face_processed_ids = set(r[0] for r in cursor.fetchall())

                cursor.execute("SELECT person_id, embedding_json FROM faces WHERE embedding_json != '[]'")
                for p_id, emb_str in cursor.fetchall():
                    if p_id not in clusters:
                        clusters[p_id] = []
                    if len(clusters[p_id]) < 15:
                        clusters[p_id].append(json.loads(emb_str))

                cursor.execute("SELECT MAX(id) FROM people")
                p_row = cursor.fetchone()
                p_count = p_row[0] if (p_row and p_row[0]) else 0

                # Build initial numpy matrix for clustering
                cluster_embs = []
                for pid, embs in clusters.items():
                    cluster_ids_list.extend([pid] * len(embs))
                    cluster_embs.extend(embs)
                if cluster_embs:
                    cm = np.array(cluster_embs)
                    cn = np.linalg.norm(cm, axis=1, keepdims=True)
                    cluster_matrix_norm = cm / np.where(cn == 0, 1, cn)

            if run_object:
                cursor.execute('''CREATE TABLE IF NOT EXISTS processed_objects (file_id INTEGER PRIMARY KEY)''')
                cursor.execute("SELECT COUNT(*) FROM processed_objects")
                if cursor.fetchone()[0] == 0:
                    with SessionLocal() as s:
                        tagged_photos = s.query(FileIndex.id).filter(FileIndex.category == 'photo', FileIndex.tags.like('%object:%')).all()
                        for (p_id,) in tagged_photos:
                            cursor.execute("INSERT OR IGNORE INTO processed_objects (file_id) VALUES (?)", (p_id,))
                cursor.execute("SELECT file_id FROM processed_objects")
                object_processed_ids = set(r[0] for r in cursor.fetchall())
            conn.commit()

        # --- Build File List ---
        files_to_process = []
        if run_index:
            backup_configs = cfg.get("backup_configs", [])
            roots = [Path(c.get("backup_path", "")) for c in backup_configs if c.get("backup_path")]
            valid_roots = [r for r in roots if r.exists() and r.is_dir()]
            for root_path in valid_roots:
                for dirpath, _, filenames in os.walk(str(root_path)):
                    for f in filenames:
                        files_to_process.append(os.path.join(dirpath, f))
        else:
            with SessionLocal() as s:
                photos = s.query(FileIndex.path).filter(FileIndex.category == 'photo').all()
                files_to_process = [p[0] for p in photos]

        total_files = len(files_to_process)
        if run_index:
            STATE["total"] = total_files
            STATE["current"] = 0
            STATE["indexed"] = 0
            STATE["status"] = "Indexing & Scanning..."
            STATE["running"] = True
        if run_face:
            STATE["face_scanner_total"] = total_files
            STATE["face_scanner_current"] = 0
        if run_object:
            STATE["object_scanner_total"] = total_files
            STATE["object_scanner_current"] = 0
            
        processed_count = 0
        
        with sqlite3.connect(ai_db_path, timeout=15) as conn:
            conn.execute("PRAGMA journal_mode=WAL;")
            cursor = conn.cursor()
            with SessionLocal() as session:
                # Store lightweight info mapping to avoid millions of session.get() queries
                file_cache = {}
                for r in session.query(FileIndex.id, FileIndex.path, FileIndex.size, FileIndex.modified, FileIndex.category, FileIndex.filename).yield_per(5000):
                    file_cache[r.path] = {"id": r.id, "size": r.size, "modified": r.modified, "category": r.category, "filename": r.filename}
                
                for idx, file_str in enumerate(files_to_process):
                    while STATE.get("paused"):
                        time.sleep(0.5)
                        if combined_scanner_stopped or (run_index and STATE.get("stopped")):
                            break
                        if not run_index and run_face and not face_scanner_running:
                            break

                    if combined_scanner_stopped or (run_index and STATE.get("stopped")):
                        break
                    if not run_index and run_face and not face_scanner_running:
                        break
                        
                    file = Path(file_str)
                    if not file.exists():
                        continue
                        
                    if run_index:
                        STATE["current"] += 1
                        STATE["current_file"] = str(file)
                    if run_face and face_scanner_running:
                        STATE["face_scanner_current"] += 1
                        STATE["face_scanner_current_file"] = file.name
                    if run_object and not STATE.get("object_scanner_stopped"):
                        STATE["object_scanner_current"] += 1
                        STATE["object_scanner_current_file"] = file.name

                    # --- 1. Indexing Phase ---
                    cached_info = file_cache.get(file_str)
                    db_item_id = cached_info["id"] if cached_info else None
                    db_item = None
                    
                    if run_index:
                        try:
                            f_stat = file.stat()
                            modified_time = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(f_stat.st_mtime))
                            file_size = str(f_stat.st_size)
                            
                            if not cached_info:
                                category = classify(file.suffix)
                                metadata, extra_tags = extract_metadata_for_file(file, category)
                                tags = build_tags(metadata, category, file.suffix, file)
                                if extra_tags: tags = ",".join(set(tags.split(",") + extra_tags))
                                
                                db_item = FileIndex(
                                    path=str(file), filename=file.name, category=category,
                                    size=file_size, modified=modified_time, extension=file.suffix,
                                    tags=tags, metadata_json=json.dumps(metadata)
                                )
                                session.add(db_item)
                                session.flush()
                                STATE["indexed"] += 1
                                if STATE["indexed"] % 500 == 0:
                                    session.commit()
                                db_item_id = db_item.id
                                file_cache[file_str] = {"id": db_item.id, "size": file_size, "modified": modified_time, "category": category, "filename": file.name}
                                cached_info = file_cache[file_str]
                            else:
                                if cached_info["size"] != file_size or cached_info["modified"] != modified_time:
                                    db_item = session.get(FileIndex, db_item_id)
                                    category = classify(file.suffix)
                                    metadata, extra_tags = extract_metadata_for_file(file, category)
                                    tags = build_tags(metadata, category, file.suffix, file)
                                    if extra_tags: tags = ",".join(set(tags.split(",") + extra_tags))
                                    db_item.size = file_size
                                    db_item.modified = modified_time
                                    db_item.metadata_json = json.dumps(metadata)
                                    db_item.tags = tags
                                    
                                    cached_info["size"] = file_size
                                    cached_info["modified"] = modified_time
                                    cached_info["category"] = category
                                    
                                    STATE["indexed"] += 1
                                    if STATE["indexed"] % 500 == 0:
                                        session.commit()
                        except Exception as e:
                            print(f"Index error on {file.name}: {e}")
                            continue

                    # Skip AI phase if the file is not an image
                    if not cached_info or cached_info["category"] != 'photo':
                        continue

                    # --- 2. AI Phase ---
                    obj_stopped = STATE.get("object_scanner_stopped", False)
                    needs_face = run_face and face_scanner_running and db_item_id not in face_processed_ids
                    needs_object = run_object and not obj_stopped and db_item_id not in object_processed_ids
                    
                    if not needs_face and not needs_object:
                        continue
                        
                    # Fetch full SQLAlchemy object ONLY if we need to modify tags or it wasn't fetched yet
                    if not db_item:
                        db_item = session.get(FileIndex, db_item_id)
                        
                    filename_lower = cached_info["filename"].lower() if cached_info["filename"] else ""
                    if "screenshot" in filename_lower or "meme" in filename_lower:
                        if needs_face: cursor.execute("INSERT OR IGNORE INTO processed_files (file_id) VALUES (?)", (db_item_id,))
                        if needs_object: cursor.execute("INSERT OR IGNORE INTO processed_objects (file_id) VALUES (?)", (db_item_id,))
                        processed_count += 1
                        if processed_count % 500 == 0:
                            conn.commit()
                            session.commit()
                        continue

                    # --- OPTIMIZATION: Read file ONCE from disk for both ML models ---
                    try:
                        img_array = np.fromfile(str(file), np.uint8)
                        img = cv2.imdecode(img_array, cv2.IMREAD_COLOR)
                    except Exception:
                        continue

                    if img is None:
                        if needs_face: cursor.execute("INSERT OR IGNORE INTO processed_files (file_id) VALUES (?)", (db_item_id,))
                        if needs_object: cursor.execute("INSERT OR IGNORE INTO processed_objects (file_id) VALUES (?)", (db_item_id,))
                        processed_count += 1
                        if processed_count % 500 == 0:
                            conn.commit()
                            session.commit()
                        continue
                        
                    # --- Run Object Classifier ---
                    if needs_object and net is not None:
                        try:
                            o_img = cv2.resize(img, (224, 224))
                            o_img = cv2.cvtColor(o_img, cv2.COLOR_BGR2RGB)
                            o_img = o_img.astype(np.float32) / 255.0
                            o_img -= np.array([0.485, 0.456, 0.406])
                            o_img /= np.array([0.229, 0.224, 0.225])
                            o_img = o_img.transpose(2, 0, 1)
                            o_img = np.expand_dims(o_img, axis=0)
                            o_img = np.ascontiguousarray(o_img)

                            net.setInput(o_img)
                            preds = net.forward().flatten()
                            exp_preds = np.exp(preds - np.max(preds))
                            probs = exp_preds / np.sum(exp_preds)
                            
                            classIds = np.argsort(probs)[-5:][::-1]
                            new_tags = []
                            for classId in classIds:
                                if probs[classId] > object_threshold:
                                    label = classes[classId].split(',')[0].strip().lower().replace(" ", "_")
                                    new_tags.append(f"object:{label}")
                                    
                            if new_tags:
                                current_tags = db_item.tags or ""
                                for tag in new_tags:
                                    if tag not in current_tags:
                                        current_tags = f"{current_tags} {tag}".strip()
                                db_item.tags = current_tags
                        except Exception as e:
                            print(f"ERROR: Failed to classify {file.name}: {e}")
                            
                        cursor.execute("INSERT OR IGNORE INTO processed_objects (file_id) VALUES (?)", (db_item_id,))
                        object_processed_ids.add(db_item_id)

                    # --- Run Face Detector ---
                    if needs_face and detector is not None:
                        try:
                            height, width, _ = img.shape
                            target_dim = 800
                            scale = 1.0
                            if max(height, width) > target_dim:
                                scale = target_dim / max(height, width)
                                new_w, new_h = int(width * scale), int(height * scale)
                                det_img = cv2.resize(img, (new_w, new_h))
                                detector.setInputSize((new_w, new_h))
                            else:
                                det_img = img
                                detector.setInputSize((width, height))
                                
                            success, faces = detector.detect(det_img)
                            
                            if faces is None:
                                target_dim = 320
                                scale = 1.0
                                if max(height, width) > target_dim:
                                    scale = target_dim / max(height, width)
                                    new_w, new_h = int(width * scale), int(height * scale)
                                    det_img = cv2.resize(img, (new_w, new_h))
                                    detector.setInputSize((new_w, new_h))
                                else:
                                    det_img = img
                                    detector.setInputSize((width, height))
                                try:
                                    success, faces = detector.detect(det_img)
                                except Exception:
                                    pass

                            if faces is not None:
                                if scale != 1.0:
                                    faces[:, :14] /= scale
                                for face in faces:
                                    face_align = recognizer.alignCrop(img, face)
                                    face_feature = recognizer.feature(face_align)
                                    embedding = face_feature[0].tolist()

                                    best_match_id = None
                                    best_sim = -1.0
                                    
                                    if cluster_matrix_norm is not None:
                                        emb_np = np.array(embedding)
                                        emb_norm = np.linalg.norm(emb_np)
                                        if emb_norm > 0:
                                            emb_np_norm = emb_np / emb_norm
                                            similarities = np.dot(cluster_matrix_norm, emb_np_norm)
                                            max_idx = np.argmax(similarities)
                                            max_sim = similarities[max_idx]
                                            
                                            if max_sim > cluster_threshold:
                                                best_sim = float(max_sim)
                                                best_match_id = cluster_ids_list[max_idx]

                                    if best_match_id is None:
                                        while True:
                                            p_count += 1
                                            cursor.execute("INSERT OR IGNORE INTO people (name) VALUES (?)", (f"Unknown Person #{p_count}",))
                                            if cursor.rowcount > 0:
                                                best_match_id = cursor.lastrowid
                                                break
                                        clusters[best_match_id] = [embedding]
                                        
                                        emb_np = np.array(embedding)
                                        emb_norm = np.linalg.norm(emb_np)
                                        emb_np_norm = emb_np / emb_norm if emb_norm > 0 else emb_np
                                        if cluster_matrix_norm is None:
                                            cluster_matrix_norm = np.array([emb_np_norm])
                                            cluster_ids_list = [best_match_id]
                                        else:
                                            cluster_matrix_norm = np.vstack([cluster_matrix_norm, emb_np_norm])
                                            cluster_ids_list.append(best_match_id)
                                    else:
                                        if len(clusters[best_match_id]) < 15:
                                            clusters[best_match_id].append(embedding)
                                            
                                            emb_np = np.array(embedding)
                                            emb_norm = np.linalg.norm(emb_np)
                                            emb_np_norm = emb_np / emb_norm if emb_norm > 0 else emb_np
                                            cluster_matrix_norm = np.vstack([cluster_matrix_norm, emb_np_norm])
                                            cluster_ids_list.append(best_match_id)
                                    cursor.execute("INSERT OR IGNORE INTO faces (person_id, file_id, embedding_json) VALUES (?, ?, ?)",
                                                    (best_match_id, db_item_id, json.dumps(embedding)))
                        except Exception as e:
                            print(f"Face processing error on {file.name}: {e}")
                            
                        cursor.execute("INSERT OR IGNORE INTO processed_files (file_id) VALUES (?)", (db_item_id,))
                        face_processed_ids.add(db_item_id)

                    processed_count += 1
                    if processed_count % 500 == 0:
                        session.commit()
                        conn.commit()

                session.commit()
                conn.commit()

    except Exception as e:
        print(f"CRITICAL: Unified Worker Error: {e}")
        traceback.print_exc()
    finally:
        if run_index:
            STATE["running"] = False
            STATE["stopped"] = False
            STATE["status"] = "Idle"
        if run_face:
            face_scanner_running = False
            STATE["face_scanner_current_file"] = ""
        if run_object:
            object_scanner_running = False
            STATE["object_scanner_stopped"] = False
            STATE["object_scanner_current_file"] = ""
        combined_scanner_running = False
        combined_scanner_stopped = False

@app.post("/scan-faces")
def scan_faces():
    if cv2 is None:
        raise HTTPException(status_code=500, detail="OpenCV is required for face recognition.")
    global face_scanner_thread, face_scanner_running
    if face_scanner_running:
        raise HTTPException(status_code=400, detail="Face scanning is already in progress.")
    face_scanner_running = True
    face_scanner_thread = threading.Thread(target=_process_unified_scanners, kwargs={"run_index": False, "run_face": True, "run_object": False})
    face_scanner_thread.start()
    return {"message": "Face scanning and clustering started in the background."}

@app.post("/stop-scan-faces")
def stop_scan_faces():
    global face_scanner_running
    if not face_scanner_running:
        raise HTTPException(status_code=400, detail="Face scanner is not running.")
    face_scanner_running = False
    return {"message": "Stopping face scanner."}

@app.get("/people")
def get_people(min_unknown_photos: int = 1):
    try:
        cfg = load_config()
        ai_db_path = get_ai_db_path()
        if not ai_db_path.exists():
            return []
        with sqlite3.connect(ai_db_path, timeout=15) as conn:
            conn.execute("PRAGMA journal_mode=WAL;")
            cursor = conn.cursor()
            cursor.execute("""
                SELECT p.id, p.name, f.file_id, COUNT(f.id) as face_count
                FROM people p
                JOIN faces f ON p.id = f.person_id
                GROUP BY p.id
                HAVING p.name NOT LIKE 'Unknown Person%' OR face_count >= ?
                ORDER BY face_count DESC
            """, (min_unknown_photos,))
            
            results = []
            for row in cursor.fetchall():
                person_id, name, sample_file_id, count = row
                results.append({
                    "id": person_id, 
                    "name": name, 
                    "face_count": count, 
                    "thumbnail": f"/people/{person_id}/thumbnail"
                })
            return results
    except Exception as e:
        print(f"Error in /people API: {e}")
        return []

@app.get("/people/{person_id}/similar-unknowns")
def get_similar_unknowns(person_id: int, threshold: float = 0.55):
    cfg = load_config()
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
         raise HTTPException(status_code=404, detail="Database not found")
         
    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        
        try:
            # Fetch known person's embeddings
            cursor.execute("SELECT id, file_id, embedding_json FROM faces WHERE person_id = ? AND embedding_json != '[]'", (person_id,))
            known_rows = cursor.fetchall()
        except sqlite3.OperationalError as e:
            if "no such table" not in str(e).lower():
                raise HTTPException(status_code=500, detail=f"Database locked or unavailable: {e}")
            raise HTTPException(status_code=404, detail="AI database tables not initialized.")
        if not known_rows:
            raise HTTPException(status_code=404, detail="Known person faces not found.")
        
        current_face_count = len(known_rows)
        cached = EXEMPLAR_CACHE.get(person_id)
        if cached and cached.get("count") == current_face_count:
            known_embeddings = cached["embeddings"]
        else:
            # Calculate Curated Reference Embeddings
            file_id_to_embs = {}
            for face_id, file_id, emb_json in known_rows:
                if file_id not in file_id_to_embs:
                    file_id_to_embs[file_id] = []
                file_id_to_embs[file_id].append(json.loads(emb_json))
                
            import numpy as np
            file_ids = list(file_id_to_embs.keys())
            
            with SessionLocal() as s:
                files_info = []
                for i in range(0, len(file_ids), 900):
                    chunk = file_ids[i:i+900]
                    chunk_info = s.query(FileIndex.id, FileIndex.modified, FileIndex.path).filter(FileIndex.id.in_(chunk)).all()
                    files_info.extend(chunk_info)
                    
            files_info.sort(key=lambda x: str(x.modified or ""))
            
            # Sample up to 50 images evenly distributed across their timeline
            sample_size = min(50, len(files_info))
            if len(files_info) > sample_size:
                indices = np.linspace(0, len(files_info)-1, sample_size, dtype=int)
                sample_files = [files_info[i] for i in indices]
            else:
                sample_files = files_info
                
            yunet_path = get_bundled_model_path("face_detection_yunet_2023mar.onnx")
            
            analyzed_files = []
            import concurrent.futures
            
            with concurrent.futures.ThreadPoolExecutor(max_workers=4) as executor:
                futures = {}
                for f_item in sample_files:
                    file_path = _resolve_path(Path(f_item.path))
                    if file_path.exists():
                        futures[executor.submit(_evaluate_image_faces, file_path, yunet_path)] = f_item
                        
                for future in concurrent.futures.as_completed(futures):
                    f_item = futures[future]
                    metrics = future.result()
                    if metrics:
                        best_metric = max(metrics, key=lambda x: x["score"])
                        analyzed_files.append({
                            "file_id": f_item.id, "date": str(f_item.modified or ""),
                            "area": best_metric["area"], "sharpness": best_metric["sharpness"]
                        })
                        
            selected_file_ids = set()
            if analyzed_files:
                analyzed_files.sort(key=lambda x: x["sharpness"], reverse=True)
                non_blurry = analyzed_files[:max(25, int(len(analyzed_files) * 0.75))]
                
                if non_blurry:
                    non_blurry.sort(key=lambda x: x["date"], reverse=True)
                    for item in non_blurry[:5]: selected_file_ids.add(item["file_id"])
                    non_blurry.sort(key=lambda x: x["date"])
                    for item in non_blurry[:5]: selected_file_ids.add(item["file_id"])
                    mid_idx = len(non_blurry) // 2
                    start_mid = max(0, mid_idx - 2)
                    for item in non_blurry[start_mid:start_mid+5]: selected_file_ids.add(item["file_id"])
                    non_blurry.sort(key=lambda x: x["area"], reverse=True)
                    for item in non_blurry[:5]: selected_file_ids.add(item["file_id"])
                    non_blurry.sort(key=lambda x: x["area"])
                    valid_far = [x for x in non_blurry if x["area"] > 0]
                    for item in valid_far[:5]: selected_file_ids.add(item["file_id"])
                    
            selected_embeddings = []
            for fid in selected_file_ids:
                if fid in file_id_to_embs:
                    selected_embeddings.extend(file_id_to_embs[fid])
                    
            if len(selected_embeddings) < 25:
                all_embs = [json.loads(row[2]) for row in known_rows if row[2]]
                for emb in all_embs:
                    if len(selected_embeddings) >= 25: break
                    if emb not in selected_embeddings: selected_embeddings.append(emb)
                        
            known_embeddings = selected_embeddings[:25]
            EXEMPLAR_CACHE.put(person_id, {"count": current_face_count, "embeddings": known_embeddings})

        # Fetch all unknown persons and their embeddings
        cursor.execute("""
            SELECT p.id, p.name, f.embedding_json, 
                   (SELECT COUNT(id) FROM faces WHERE person_id = p.id) as photo_count,
                   f.file_id
            FROM people p
            JOIN faces f ON p.id = f.person_id
            WHERE p.name LIKE 'Unknown Person%' AND f.embedding_json != '[]'
        """)
        
        try:
            import numpy as np
            has_numpy = True
            known_matrix = np.array(known_embeddings)
            known_norms = np.linalg.norm(known_matrix, axis=1, keepdims=True)
            known_matrix_norm = known_matrix / np.where(known_norms == 0, 1, known_norms)
        except ImportError:
            has_numpy = False
        
        similar_profiles = {}
        for unk_person_id, unk_name, unk_embedding_json, photo_count, file_id in cursor:
            if not unk_embedding_json:
                continue
            
            if has_numpy:
                unk_embedding = np.array(json.loads(unk_embedding_json))
                unk_norm = np.linalg.norm(unk_embedding)
                if unk_norm == 0:
                    continue
                unk_embedding_norm = unk_embedding / unk_norm
                similarities = np.dot(known_matrix_norm, unk_embedding_norm)
                max_sim = np.max(similarities)
            else:
                unk_embedding = json.loads(unk_embedding_json)
                max_sim = 0.0
                for known_emb in known_embeddings:
                    if known_emb:
                        sim = _cosine_similarity(known_emb, unk_embedding)
                        if sim > max_sim:
                            max_sim = sim
                            
            if max_sim >= threshold:
                if unk_person_id not in similar_profiles or max_sim > similar_profiles[unk_person_id]["similarity"]:
                    similar_profiles[unk_person_id] = {
                        "id": unk_person_id, "name": unk_name, "similarity": round(float(max_sim), 3),
                        "face_count": photo_count, "thumbnail": f"/people/{unk_person_id}/thumbnail?v={file_id}_{photo_count}"
                    }
        results = list(similar_profiles.values())
        
        with SessionLocal() as s:
            for match in results:
                # Get a sample of file IDs for this unknown person
                cursor.execute("SELECT file_id FROM faces WHERE person_id = ? LIMIT 10", (match["id"],))
                file_ids = [r[0] for r in cursor.fetchall()]
                
                if file_ids:
                    # Look up their paths and dates in the main database
                    files_info = s.query(FileIndex.path, FileIndex.modified).filter(FileIndex.id.in_(file_ids)).all()
                    match["sample_paths"] = "|".join([str(f.path) for f in files_info if f.path])
                    match["sample_dates"] = "|".join([str(f.modified) for f in files_info if f.modified])
                    
        results.sort(key=lambda x: x["similarity"], reverse=True)
        return results

@app.get("/people/{person_id}/thumbnail")
def get_person_thumbnail(person_id: int, theme: str = "dark"):
    if cv2 is None:
        raise HTTPException(status_code=500, detail="OpenCV not installed")
        
    cfg = load_config()
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
        raise HTTPException(status_code=404, detail="Database not found")

    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        cursor.execute("SELECT thumbnail_file_id FROM people WHERE id = ?", (person_id,))
        thumb_row = cursor.fetchone()
        thumb_file_id = thumb_row[0] if thumb_row else None
            
        face_row = None
        if thumb_file_id:
            cursor.execute("SELECT file_id, embedding_json FROM faces WHERE person_id = ? AND file_id = ? AND embedding_json != '[]' LIMIT 1", (person_id, thumb_file_id))
            face_row = cursor.fetchone()
            if not face_row:
                # User selected a manual tag as cover photo (no embedding). Return the full photo fallback.
                return preview(thumb_file_id)
            
        if not face_row:
            cursor.execute("SELECT file_id, embedding_json FROM faces WHERE person_id = ? AND embedding_json != '[]' ORDER BY id DESC LIMIT 1", (person_id,))
            face_row = cursor.fetchone()
            
        if not face_row:
            # If they only have manual tags without embeddings, fallback to the full uncropped photo
            cursor.execute("SELECT file_id FROM faces WHERE person_id = ? ORDER BY id DESC LIMIT 1", (person_id,))
            fallback = cursor.fetchone()
            if fallback:
                return preview(fallback[0], theme)
            raise HTTPException(status_code=404, detail="Person not found")
        file_id, emb_json = face_row
        target_embedding = json.loads(emb_json)

    thumb_dir = Path(cfg.get("thumbnail_path") or "thumbnails") / ".wabs_cache" / "faces"
    thumb_dir.mkdir(parents=True, exist_ok=True)
    cached_face = thumb_dir / f"person_{person_id}.jpg"
    
    if cached_face.exists():
        return FileResponse(str(cached_face), media_type="image/jpeg")

    try:
        with SessionLocal() as s:
            file_item = s.query(FileIndex).filter(FileIndex.id == file_id).first()
            if not file_item:
                raise HTTPException(status_code=404, detail="Image not found")
        file_path = _resolve_path(Path(file_item.path))
        
        import numpy as np
        img_array = np.fromfile(str(file_path), np.uint8)
        img = cv2.imdecode(img_array, cv2.IMREAD_COLOR)
        
        if img is None:
            try:
                from PIL import Image, ImageOps
                with Image.open(file_path) as pil_img:
                    pil_img = ImageOps.exif_transpose(pil_img)
                    if pil_img.mode != 'RGB':
                        pil_img = pil_img.convert('RGB')
                    img = cv2.cvtColor(np.array(pil_img), cv2.COLOR_RGB2BGR)
            except Exception as e:
                print(f"Pillow face thumbnail fallback failed for {file_path.name}: {e}")
                
        if img is None:
            return preview(file_id, theme)

        yunet_path = get_bundled_model_path("face_detection_yunet_2023mar.onnx")
        sface_path = get_bundled_model_path("face_recognition_sface_2021dec.onnx")

        if not Path(yunet_path).exists() or not Path(sface_path).exists():
            print("Face recognition models not found. Ensure .onnx files are in the backend folder.")
            return

        detector = cv2.FaceDetectorYN.create(yunet_path, "", (320, 320))
        detector.setScoreThreshold(0.5)
        recognizer = cv2.FaceRecognizerSF.create(sface_path, "")

        height, width, _ = img.shape
        target_dim = 800
        scale = 1.0
        if max(height, width) > target_dim:
            scale = target_dim / max(height, width)
            new_w, new_h = int(width * scale), int(height * scale)
            det_img = cv2.resize(img, (new_w, new_h))
            detector.setInputSize((new_w, new_h))
        else:
            det_img = img
            detector.setInputSize((width, height))

        success, faces = detector.detect(det_img)

        if faces is None:
            target_dim = 320
            scale = 1.0
            if max(height, width) > target_dim:
                scale = target_dim / max(height, width)
                new_w, new_h = int(width * scale), int(height * scale)
                det_img = cv2.resize(img, (new_w, new_h))
                detector.setInputSize((new_w, new_h))
            else:
                det_img = img
                detector.setInputSize((width, height))
            success, faces = detector.detect(det_img)

        if faces is not None:
            if scale != 1.0:
                faces[:, :14] /= scale

            best_face_align = None
            best_sim = -1.0
            
            target_emb_np = np.array(target_embedding)
            target_norm = np.linalg.norm(target_emb_np)
            target_emb_norm = target_emb_np / target_norm if target_norm > 0 else target_emb_np
            
            for face in faces:
                face_align = recognizer.alignCrop(img, face)
                face_feature = recognizer.feature(face_align)
                
                emb_np = face_feature[0]
                emb_norm = np.linalg.norm(emb_np)
                emb_np_norm = emb_np / emb_norm if emb_norm > 0 else emb_np
                
                sim = float(np.dot(target_emb_norm, emb_np_norm))
                
                if sim > best_sim:
                    best_sim = sim
                    best_face_align = face_align
                    
                # Early Exit: Since the target embedding came from this exact image, 
                # the match will be nearly 1.0. We can safely skip the remaining faces!
                if best_sim > 0.98:
                    break

            if best_face_align is not None and best_sim >= 0.40:
                is_success, buffer = cv2.imencode(".jpg", best_face_align)
                if is_success:
                    with open(str(cached_face), "wb") as f:
                        f.write(buffer.tobytes())
                if cached_face.exists():
                    return FileResponse(str(cached_face), media_type="image/jpeg")

    except Exception as e:
        print(f"Failed to generate face thumbnail: {e}")

    # Fallback to the full image thumbnail if face crop fails
    return preview(file_id, theme)

@app.get("/people/{person_id}/photos")
def get_person_photos(person_id: int, offset: int = 0, limit: int = 50):
    cfg = load_config()
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
        return []
        
    ui_prefs = cfg.get("ui_preferences") or {}
    cache_enabled = cfg.get("enable_photo_thumbnail_cache")
    if cache_enabled is None:
        cache_enabled = ui_prefs.get("enable_photo_thumbnail_cache", False)
    cache_flag = "&tc=1" if str(cache_enabled).lower() in ("true", "1", "yes") else ""

    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        try:
            cursor.execute(
                "SELECT file_id FROM faces WHERE person_id = ? GROUP BY file_id ORDER BY file_id DESC LIMIT ? OFFSET ?", 
                (person_id, limit, offset)
            )
            file_ids = [r[0] for r in cursor.fetchall()]
        except sqlite3.OperationalError as e:
            if "no such table" not in str(e).lower():
                raise HTTPException(status_code=500, detail=f"Database locked or unavailable: {e}")
            return []
    if not file_ids:
        return []
    with SessionLocal() as s:
        results = []
        # Chunk queries to prevent SQLite IN() limitations and memory exhaustion
        for i in range(0, len(file_ids), 900):
            chunk = file_ids[i:i + 900]
            photos = s.query(FileIndex.id, FileIndex.filename, FileIndex.path, FileIndex.category, FileIndex.size, FileIndex.modified, FileIndex.extension, FileIndex.tags, FileIndex.metadata_json).filter(FileIndex.id.in_(chunk)).all()
            
            # Ensure the response maintains the exact ordered pagination from SQLite
            photo_dict = {p.id: _build_item(p, cache_flag) for p in photos}
            results.extend([photo_dict[fid] for fid in chunk if fid in photo_dict])
        return results

@app.post("/people/{person_id}/set-thumbnail")
def set_person_thumbnail(person_id: int, payload: dict = Body(...)):
    file_id = payload.get("file_id")
    if not file_id:
        raise HTTPException(status_code=400, detail="file_id is required")

    # Invalidate cache for this person
    EXEMPLAR_CACHE.pop(person_id)

    cfg = load_config()
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
         raise HTTPException(status_code=404, detail="Database not found")
         
    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        cursor.execute("UPDATE people SET thumbnail_file_id = ? WHERE id = ?", (file_id, person_id))
        conn.commit()
        
    thumb_dir = Path(cfg.get("thumbnail_path") or "thumbnails") / ".wabs_cache" / "faces"
    cached_face = thumb_dir / f"person_{person_id}.jpg"
    if cached_face.exists():
        try:
            cached_face.unlink()
        except Exception:
            pass
            
    return {"success": True}

@app.post("/people/{person_id}/suggest-thumbnail")
def auto_suggest_thumbnail(person_id: int):
    if cv2 is None:
        raise HTTPException(status_code=500, detail="OpenCV not installed")
        
    cfg = load_config()
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
         raise HTTPException(status_code=404, detail="Database not found")

    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        
        # Fetch up to 30 recent faces to evaluate (to avoid locking the CPU for too long)
        cursor.execute("""
            SELECT file_id 
            FROM faces 
            WHERE person_id = ? AND embedding_json != '[]' 
            ORDER BY id DESC LIMIT 30
        """, (person_id,))
        face_rows = cursor.fetchall()

    if not face_rows:
        raise HTTPException(status_code=404, detail="No valid faces found for this person.")

    yunet_path = get_bundled_model_path("face_detection_yunet_2023mar.onnx")

    best_file_id = None
    best_score = -1.0

    import concurrent.futures
    
    file_items = []
    with SessionLocal() as s:
        for (file_id,) in face_rows:
            file_item = s.query(FileIndex).filter(FileIndex.id == file_id).first()
            if file_item:
                file_items.append((file_id, _resolve_path(Path(file_item.path))))
                
    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as executor:
        futures = {executor.submit(_evaluate_image_faces, fp, yunet_path): fid for fid, fp in file_items if fp.exists()}
        for future in concurrent.futures.as_completed(futures):
            fid = futures[future]
            metrics = future.result()
            for fm in metrics:
                if fm["score"] > best_score:
                    best_score = fm["score"]
                    best_file_id = fid

    if best_file_id:
        # Save the best thumbnail back to the database
        with sqlite3.connect(ai_db_path, timeout=15) as conn:
            cursor = conn.cursor()
            cursor.execute("UPDATE people SET thumbnail_file_id = ? WHERE id = ?", (best_file_id, person_id))
            conn.commit()
            
        # Clear the old cached thumbnail so it regenerates on next load
        thumb_dir = Path(cfg.get("thumbnail_path") or "thumbnails") / ".wabs_cache" / "faces"
        cached_face = thumb_dir / f"person_{person_id}.jpg"
        if cached_face.exists():
            try:
                cached_face.unlink()
            except Exception:
                pass
                
        if cfg.get("enable_logging"):
            import logging
            logging.info(f"Auto-suggested cover photo {best_file_id} for person {person_id} with score {round(best_score, 2)}")
            
        return {"success": True, "new_thumbnail_id": best_file_id, "score": best_score}
        
    raise HTTPException(status_code=400, detail="Could not determine a suitable thumbnail.")

@app.post("/people/{person_id}/remove-photo")
def remove_person_photo(person_id: int, payload: dict = Body(...)):
    file_id = payload.get("file_id")
    if not file_id:
        raise HTTPException(status_code=400, detail="file_id is required")

    # Invalidate cache for this person
    EXEMPLAR_CACHE.pop(person_id)

    cfg = load_config()
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
         raise HTTPException(status_code=404, detail="Database not found")

    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM people WHERE id = ?", (person_id,))
        person_row = cursor.fetchone()
        person_name = person_row[0] if person_row else None
        
        cursor.execute("DELETE FROM faces WHERE person_id = ? AND file_id = ?", (person_id, file_id))
        deleted_count = cursor.rowcount
        
        cursor.execute("SELECT thumbnail_file_id FROM people WHERE id = ?", (person_id,))
        thumb_row = cursor.fetchone()
        if thumb_row and thumb_row[0] == file_id:
            cursor.execute("UPDATE people SET thumbnail_file_id = NULL WHERE id = ?", (person_id,))
            
        conn.commit()
        
    if deleted_count > 0:
        if person_name and not person_name.startswith("Unknown Person"):
            with SessionLocal() as s:
                f = s.query(FileIndex).filter(FileIndex.id == file_id).first()
                if f and f.tags:
                        current_tags = set(f.tags.split())
                        tag_to_remove = f"person:{person_name}"
                        if tag_to_remove in current_tags:
                            current_tags.remove(tag_to_remove)
                            f.tags = " ".join(sorted(current_tags))
                            s.commit()

        thumb_dir = Path(cfg.get("thumbnail_path") or "thumbnails") / ".wabs_cache" / "faces"
        cached_face = thumb_dir / f"person_{person_id}.jpg"
        if cached_face.exists():
            try:
                cached_face.unlink()
            except Exception:
                pass

    return {"success": True, "removed": deleted_count}

@app.post("/people/{person_id}/add-photo")
def add_person_photo(person_id: int, payload: dict = Body(...)):
    file_id = payload.get("file_id")
    if not file_id:
        raise HTTPException(status_code=400, detail="file_id is required")

    cfg = load_config()
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
         raise HTTPException(status_code=404, detail="Database not found")

    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM people WHERE id = ?", (person_id,))
        person_row = cursor.fetchone()
        if not person_row:
            raise HTTPException(status_code=404, detail="Person not found")
        person_name = person_row[0]
        
        cursor.execute("SELECT id FROM faces WHERE person_id = ? AND file_id = ?", (person_id, file_id))
        if cursor.fetchone():
            return {"success": True, "message": "Already tagged"}
            
        # Insert empty array since it's a manual tag (bypasses similarity checks)
        cursor.execute("INSERT OR IGNORE INTO faces (person_id, file_id, embedding_json) VALUES (?, ?, ?)", (person_id, file_id, "[]"))
        conn.commit()
        
    if person_name and not person_name.startswith("Unknown Person"):
        with SessionLocal() as s:
            f = s.query(FileIndex).filter(FileIndex.id == file_id).first()
            if f:
                current_tags = set((f.tags or "").split())
                new_tag = f"person:{person_name}"
                if new_tag not in current_tags:
                    current_tags.add(new_tag)
                    f.tags = " ".join(sorted(current_tags))
                    s.commit()

    return {"success": True}

@app.post("/people/{person_id}/rename")
def rename_person(person_id: int, payload: dict = Body(...)):
    new_name = payload.get("name", "Unknown Person").strip()
    if not new_name:
        new_name = "Unknown Person"

    cfg = load_config()
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
         raise HTTPException(status_code=404, detail="Database not found")
         
    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM people WHERE id = ?", (person_id,))
        old_name_row = cursor.fetchone()
        old_name = old_name_row[0] if old_name_row else "Unknown Person"
        
        cursor.execute("SELECT DISTINCT file_id FROM faces WHERE person_id = ?", (person_id,))
        file_ids = [r[0] for r in cursor.fetchall()]
        
        # Case-insensitive check to see if the target name already exists
        cursor.execute("SELECT id FROM people WHERE name COLLATE NOCASE = ? AND id != ?", (new_name, person_id))
        existing_person = cursor.fetchone()
        
        if existing_person:
            target_id = existing_person[0]
            # Invalidate both caches if a merge happens
            EXEMPLAR_CACHE.pop(target_id)
            # Auto-Merge: Reassign all faces to the existing person, then delete the duplicate
            cursor.execute("UPDATE OR IGNORE faces SET person_id = ? WHERE person_id = ?", (target_id, person_id))
            cursor.execute("DELETE FROM faces WHERE person_id = ?", (person_id,))
            cursor.execute("DELETE FROM people WHERE id = ?", (person_id,))
        else:
            # Standard Rename
            cursor.execute("UPDATE people SET name = ? WHERE id = ?", (new_name, person_id))
            
        EXEMPLAR_CACHE.pop(person_id)
        conn.commit()
        
    if file_ids:
        with SessionLocal() as s:
            for i in range(0, len(file_ids), 900):
                chunk = file_ids[i:i + 900]
                files_to_update = s.query(FileIndex.id, FileIndex.tags).filter(FileIndex.id.in_(chunk)).all()
                mappings = []
                for f_id, tags in files_to_update:
                    current_tags_set = set((tags or "").split())
                    if old_name and not old_name.startswith("Unknown Person"):
                        current_tags_set.discard(f"person:{old_name}")
                    if new_name and not new_name.startswith("Unknown Person"):
                        current_tags_set.add(f"person:{new_name}")
                        
                    new_tags_str = " ".join(sorted(current_tags_set))
                    if new_tags_str != tags:
                        mappings.append({"id": f_id, "tags": new_tags_str})
                if mappings:
                    s.bulk_update_mappings(FileIndex, mappings)
                    s.commit()
            
    return {"success": True, "name": new_name}

@app.delete("/people/{person_id}")
def delete_person(person_id: int):
    cfg = load_config()
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
         raise HTTPException(status_code=404, detail="Database not found")
         
    EXEMPLAR_CACHE.pop(person_id)

    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM people WHERE id = ?", (person_id,))
        old_name_row = cursor.fetchone()
        old_name = old_name_row[0] if old_name_row else "Unknown Person"
        
        cursor.execute("SELECT DISTINCT file_id FROM faces WHERE person_id = ?", (person_id,))
        file_ids = [r[0] for r in cursor.fetchall()]
        
        # Wipe the face embeddings and the person
        cursor.execute("DELETE FROM faces WHERE person_id = ?", (person_id,))
        cursor.execute("DELETE FROM people WHERE id = ?", (person_id,))
        conn.commit()
        
    # Clean up any tags from the main index
    if file_ids and old_name and not old_name.startswith("Unknown Person"):
        with SessionLocal() as s:
            for i in range(0, len(file_ids), 900):
                chunk = file_ids[i:i + 900]
                files_to_update = s.query(FileIndex.id, FileIndex.tags).filter(FileIndex.id.in_(chunk)).all()
                mappings = []
                for f_id, tags in files_to_update:
                    if tags:
                        current_tags_set = set(tags.split())
                        current_tags_set.discard(f"person:{old_name}")
                        new_tags_str = " ".join(sorted(current_tags_set))
                        if new_tags_str != tags:
                            mappings.append({"id": f_id, "tags": new_tags_str})
                if mappings:
                    s.bulk_update_mappings(FileIndex, mappings)
                    s.commit()
            
    return {"success": True, "deleted_id": person_id}

@app.post("/people/merge")
def merge_people(payload: dict = Body(...)):
    person_ids = payload.get("person_ids", [])
    if not person_ids or len(person_ids) < 2:
        raise HTTPException(status_code=400, detail="At least two person IDs are required for merging.")
    
    cfg = load_config()
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
         raise HTTPException(status_code=404, detail="Database not found")
         
    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        
        placeholders = ",".join("?" * len(person_ids))
        cursor.execute(f"SELECT id, name FROM people WHERE id IN ({placeholders})", person_ids)
        people_rows = cursor.fetchall()
        
        if not people_rows:
             raise HTTPException(status_code=404, detail="People not found.")
             
        people_rows.sort(key=lambda p: (0 if p[1] and not p[1].startswith("Unknown Person") else 1, p[0]))
        primary_id = people_rows[0][0]
        ids_to_merge = [p[0] for p in people_rows if p[0] != primary_id]
        
        for old_id in ids_to_merge:
            cursor.execute("UPDATE OR IGNORE faces SET person_id = ? WHERE person_id = ?", (primary_id, old_id))
            cursor.execute("DELETE FROM faces WHERE person_id = ?", (old_id,))
            cursor.execute("DELETE FROM people WHERE id = ?", (old_id,))
            EXEMPLAR_CACHE.pop(old_id)
            
        EXEMPLAR_CACHE.pop(primary_id)
        conn.commit()
        
    if cfg.get("enable_logging"):
        import logging
        logging.info(f"Merged {len(ids_to_merge)} unknown profiles into person {primary_id}.")
        
    return {"success": True, "merged_into": primary_id}

@app.post("/people/cluster-unknowns")
def cluster_unknowns(payload: dict = Body(...)):
    person_ids = payload.get("person_ids", [])
    threshold = payload.get("threshold", 0.55)
    if not person_ids:
        return {"merged_count": 0, "results": []}
        
    cfg = load_config()
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
         raise HTTPException(status_code=404, detail="Database not found")
         
    with sqlite3.connect(ai_db_path, timeout=60) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        
        # Fetch ONLY Unknown People embeddings to cluster them together
        cursor.execute("""
            SELECT p.id, p.name, f.embedding_json
            FROM people p
            JOIN faces f ON p.id = f.person_id
            WHERE p.name LIKE 'Unknown Person%' AND f.embedding_json != '[]'
        """)
        
        all_rows = cursor.fetchall()
        if not all_rows:
            return {"merged_count": 0, "message": "No unknown persons to compare against.", "results": []}
            
        import numpy as np
        
        person_embs = {}
        person_names = {}
        
        # OOM PREVENTION: Cap the maximum number of embeddings per cluster.
        # Unknown profiles are highly cohesive. 15 faces perfectly represent the cluster's variations
        # and completely prevent massive O(N^2) memory spikes during the matrix dot product.
        MAX_EMBS_PER_PERSON = 15
        
        for pid, pname, emb_json in all_rows:
            if pid not in person_embs:
                person_embs[pid] = []
            if len(person_embs[pid]) < MAX_EMBS_PER_PERSON:
                person_embs[pid].append(json.loads(emb_json))
            person_names[pid] = pname
            
        # Sort all PIDs by number of faces DESC so smaller clusters merge into larger ones
        sorted_pids = sorted(person_embs.keys(), key=lambda k: len(person_embs[k]), reverse=True)
        
        canonical_embs = []
        canonical_pids = []
        for pid in sorted_pids:
            canonical_embs.extend(person_embs[pid])
            canonical_pids.extend([pid] * len(person_embs[pid]))
            
        # Cast to float32 to instantly halve the memory requirements of the matrix
        k_matrix = np.array(canonical_embs, dtype=np.float32)
        k_norms = np.linalg.norm(k_matrix, axis=1, keepdims=True)
        k_matrix_norm = k_matrix / np.where(k_norms == 0, 1, k_norms)
        
        merged_count = 0
        results = []
        merged_away = set() # Track already-merged IDs so they aren't processed twice
        
        # Order target IDs ascending by size so small cards fold into large ones
        target_ids = sorted([pid for pid in person_ids if pid in person_embs], key=lambda k: len(person_embs[k]))
        
        # Batch targets to prevent out-of-memory errors on massive 30,000x30,000 matrix operations
        batch_size_embs = 250
        batches, current_batch_pids = [], []
        current_batch_embs_count = 0
        
        for pid in target_ids:
            num_embs = len(person_embs[pid])
            if current_batch_embs_count + num_embs > batch_size_embs and current_batch_pids:
                batches.append(current_batch_pids)
                current_batch_pids = []
                current_batch_embs_count = 0
            current_batch_pids.append(pid)
            current_batch_embs_count += num_embs
        if current_batch_pids:
            batches.append(current_batch_pids)
            
        for batch_pids in batches:
            if APP_SHUTTING_DOWN:
                break
            valid_batch_pids = [pid for pid in batch_pids if pid not in merged_away]
            if not valid_batch_pids:
                continue
                
            batch_embs = []
            for pid in valid_batch_pids:
                batch_embs.extend(person_embs[pid])
                
            unk_matrix = np.array(batch_embs, dtype=np.float32)
            unk_norms = np.linalg.norm(unk_matrix, axis=1, keepdims=True)
            unk_matrix_norm = unk_matrix / np.where(unk_norms == 0, 1, unk_norms)

            similarities = np.dot(k_matrix_norm, unk_matrix_norm.T)
            
            col_offset = 0
            for pid in valid_batch_pids:
                if pid in merged_away:
                    col_offset += len(person_embs[pid])
                    continue
                    
                num_embs = len(person_embs[pid])
                pid_similarities = similarities[:, col_offset:col_offset+num_embs]
                col_offset += num_embs
                
                max_sims_per_canonical = np.max(pid_similarities, axis=1)
                sorted_indices = np.argsort(max_sims_per_canonical)[::-1]
                
                best_match_id = None
                max_sim = 0.0
                
                for idx in sorted_indices:
                    match_pid = canonical_pids[idx]
                    sim = float(max_sims_per_canonical[idx])
                    if sim < threshold:
                        break # Sorted DESC, so anything below is a guaranteed fail
                        
                    if match_pid != pid and match_pid not in merged_away:
                        best_match_id = match_pid
                        max_sim = sim
                        break
                        
                if best_match_id is not None:
                    cursor.execute("UPDATE OR IGNORE faces SET person_id = ? WHERE person_id = ?", (best_match_id, pid))
                    cursor.execute("DELETE FROM faces WHERE person_id = ?", (pid,))
                    cursor.execute("DELETE FROM people WHERE id = ?", (pid,))
                    EXEMPLAR_CACHE.pop(pid)
                    EXEMPLAR_CACHE.pop(best_match_id)
                    merged_away.add(pid)
                    merged_count += 1
                    results.append({
                        "id": pid, "merged_into": best_match_id, 
                        "name": person_names[best_match_id], "similarity": round(max_sim, 3)
                    })
                    
        conn.commit()
        
    if cfg.get("enable_logging"):
        import logging
        logging.info(f"Clustered {merged_count} unknown profiles together.")
        
    return {"merged_count": merged_count, "results": results}

@app.post("/people/reclassify")
def reclassify_people(payload: dict = Body(...)):
    person_ids = payload.get("person_ids", [])
    threshold = payload.get("threshold", 0.55)
    
    if not person_ids:
        return {"reclassified_count": 0}
        
    cfg = load_config()
    ai_db_path = get_ai_db_path()
    thumb_dir = Path(cfg.get("thumbnail_path") or "thumbnails") / ".wabs_cache" / "faces"
    if not ai_db_path.exists():
         raise HTTPException(status_code=404, detail="Database not found")
         
    with sqlite3.connect(ai_db_path, timeout=60) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        
        # 1. Fetch all embeddings for the target unknown profiles in safe chunks
        faces_to_reclassify = []
        for i in range(0, len(person_ids), 900):
            chunk = person_ids[i:i+900]
            placeholders = ",".join("?" * len(chunk))
            cursor.execute(f"""
                SELECT f.id, f.file_id, f.embedding_json, p.id
                FROM faces f
                JOIN people p ON p.id = f.person_id
                WHERE p.id IN ({placeholders}) AND p.name LIKE 'Unknown Person%' AND f.embedding_json != '[]'
            """, chunk)
            faces_to_reclassify.extend(cursor.fetchall())
            
        if not faces_to_reclassify:
            return {"reclassified_count": 0, "message": "No valid faces to reclassify."}
            
        target_person_ids = list(set([r[3] for r in faces_to_reclassify]))
        
        # 2. Fetch the exemplar embeddings for all OTHER profiles (Named and Unselected Unknowns)
        cursor.execute("""
            SELECT p.id, f.embedding_json
            FROM people p
            JOIN faces f ON p.id = f.person_id
            WHERE f.embedding_json != '[]'
        """)
        
        target_ids_set = set(person_ids)
        clusters = {}
        for pid, emb_json in cursor.fetchall():
            if pid in target_ids_set:
                continue
            if pid not in clusters:
                clusters[pid] = []
            if len(clusters[pid]) < 15:
                clusters[pid].append(json.loads(emb_json))
                
        # 3. Delete the old targeted profiles and their faces
        for pid in target_person_ids:
            cursor.execute("DELETE FROM faces WHERE person_id = ?", (pid,))
            cursor.execute("DELETE FROM people WHERE id = ?", (pid,))
            EXEMPLAR_CACHE.pop(pid)
            old_thumb = thumb_dir / f"person_{pid}.jpg"
            if old_thumb.exists():
                try: old_thumb.unlink()
                except Exception: pass
            
        # Determine the current max person ID for generating new Unknowns
        cursor.execute("SELECT MAX(id) FROM people")
        p_row = cursor.fetchone()
        p_count = p_row[0] if (p_row and p_row[0]) else 0
        
        import numpy as np
        
        cluster_embs = []
        cluster_ids_list = []
        for pid, embs in clusters.items():
            cluster_ids_list.extend([pid] * len(embs))
            cluster_embs.extend(embs)
            
        if cluster_embs:
            cm = np.array(cluster_embs, dtype=np.float32)
            cn = np.linalg.norm(cm, axis=1, keepdims=True)
            cluster_matrix_norm = cm / np.where(cn == 0, 1, cn)
        else:
            cluster_matrix_norm = None
            
        # 4. Re-cluster the faces
        reclassified_count = 0
        files_to_tag = {}
        
        cursor.execute("SELECT id, name FROM people WHERE name NOT LIKE 'Unknown Person%'")
        named_people_map = {r[0]: r[1] for r in cursor.fetchall()}
        
        for face_id, file_id, emb_json, old_pid in faces_to_reclassify:
            if APP_SHUTTING_DOWN:
                break
            embedding = json.loads(emb_json)
            emb_np = np.array(embedding, dtype=np.float32)
            emb_norm = np.linalg.norm(emb_np)
            emb_np_norm = emb_np / emb_norm if emb_norm > 0 else emb_np
            
            best_match_id = None
            
            if cluster_matrix_norm is not None:
                similarities = np.dot(cluster_matrix_norm, emb_np_norm)
                max_idx = np.argmax(similarities)
                max_sim = similarities[max_idx]
                
                if max_sim >= threshold:
                    best_match_id = cluster_ids_list[max_idx]
                    
            if best_match_id is None:
                while True:
                    p_count += 1
                    cursor.execute("INSERT OR IGNORE INTO people (name) VALUES (?)", (f"Unknown Person #{p_count}",))
                    if cursor.rowcount > 0:
                        best_match_id = cursor.lastrowid
                        break
                clusters[best_match_id] = [embedding]
                if cluster_matrix_norm is None:
                    cluster_matrix_norm = np.array([emb_np_norm])
                    cluster_ids_list = [best_match_id]
                else:
                    cluster_matrix_norm = np.vstack([cluster_matrix_norm, emb_np_norm])
                    cluster_ids_list.append(best_match_id)
            else:
                if len(clusters[best_match_id]) < 15:
                    clusters[best_match_id].append(embedding)
                    cluster_matrix_norm = np.vstack([cluster_matrix_norm, emb_np_norm])
                    cluster_ids_list.append(best_match_id)
                    
            cursor.execute("INSERT OR IGNORE INTO faces (person_id, file_id, embedding_json) VALUES (?, ?, ?)",
                            (best_match_id, file_id, emb_json))
            reclassified_count += 1
            
            if best_match_id in named_people_map:
                name = named_people_map[best_match_id]
                if name not in files_to_tag:
                    files_to_tag[name] = set()
                files_to_tag[name].add(file_id)
                
        conn.commit()
        
    if files_to_tag:
        with SessionLocal() as s:
            for name, f_ids in files_to_tag.items():
                f_ids_list = list(f_ids)
                for i in range(0, len(f_ids_list), 900):
                    chunk = f_ids_list[i:i + 900]
                    files_to_update = s.query(FileIndex.id, FileIndex.tags).filter(FileIndex.id.in_(chunk)).all()
                    mappings = []
                    for f_id, tags in files_to_update:
                        current_tags_set = set((tags or "").split())
                        new_tag = f"person:{name}"
                        if new_tag not in current_tags_set:
                            current_tags_set.add(new_tag)
                            mappings.append({"id": f_id, "tags": " ".join(sorted(current_tags_set))})
                    if mappings:
                        s.bulk_update_mappings(FileIndex, mappings)
                        s.commit()
                        
    if cfg.get("enable_logging"):
        import logging
        logging.info(f"Reclassified {reclassified_count} faces from unknown profiles.")
                        
    return {"reclassified_count": reclassified_count}

object_scanner_running = False
object_scanner_thread = None
object_scanner_current = 0
object_scanner_total = 0
object_scanner_stopped = False
object_scanner_current_file = ""

def _scan_and_tag_objects_worker():
    global object_scanner_running, object_scanner_current, object_scanner_total, object_scanner_stopped, object_scanner_current_file
    try:
        import numpy as np
        # model_path = get_bundled_model_path("mobilenetv2-small.onnx")
        model_path = get_bundled_model_path("mobilenetv2-small.onnx")
        classes_path = get_bundled_model_path("imagenet_classes.txt")
        
        if not Path(model_path).exists() or not Path(classes_path).exists():
            print("Object classification models not found. Ensure mobilenetv2-small.onnx and imagenet_classes.txt are in the backend folder.")
            return

        net = cv2.dnn.readNetFromONNX(model_path)
        with open(classes_path, 'rt') as f:
            classes = [line.strip() for line in f.readlines()]
            
        cfg = load_config()
        object_sensitivity = cfg.get("object_sensitivity", "medium")
        if object_sensitivity == "high":
            object_threshold = 0.10
        elif object_sensitivity == "low":
            object_threshold = 0.30
        else:
            object_threshold = 0.15
            
        ai_db_path = get_ai_db_path()
        with sqlite3.connect(ai_db_path) as conn:
            cursor = conn.cursor()
            cursor.execute('''CREATE TABLE IF NOT EXISTS processed_objects (
                                file_id INTEGER PRIMARY KEY
                              )''')
            
            with SessionLocal() as s:
                # Seed existing tagged files to avoid rescanning legacy data
                cursor.execute("SELECT COUNT(*) FROM processed_objects")
                if cursor.fetchone()[0] == 0:
                    tagged_photos = s.query(FileIndex.id).filter(FileIndex.category == 'photo', FileIndex.tags.like('%object:%')).all()
                    for (p_id,) in tagged_photos:
                        cursor.execute("INSERT OR IGNORE INTO processed_objects (file_id) VALUES (?)", (p_id,))
                    conn.commit()

                cursor.execute("SELECT file_id FROM processed_objects")
                processed_ids = set(r[0] for r in cursor.fetchall())

                # Fetch IDs instead of all objects to save memory on large datasets
                photo_ids = [r[0] for r in s.query(FileIndex.id).filter(FileIndex.category == 'photo').all()]
                
                STATE["object_scanner_total"] = len(photo_ids)
                STATE["object_scanner_current"] = 0
                processed_count = 0

                for photo_id in photo_ids:
                    if not object_scanner_running:
                        break
                        
                    photo = s.get(FileIndex, photo_id)
                    if not photo:
                        continue
                        
                    STATE["object_scanner_current"] += 1
                    STATE["object_scanner_current_file"] = photo.filename or "Unknown file"
                    
                    if photo.id in processed_ids:
                        continue
                        
                    filename_lower = photo.filename.lower() if photo.filename else ""
                    if "screenshot" in filename_lower or "meme" in filename_lower:
                        cursor.execute("INSERT OR IGNORE INTO processed_objects (file_id) VALUES (?)", (photo.id,))
                        processed_count += 1
                        if processed_count % 500 == 0:
                            conn.commit()
                            s.commit()
                        continue

                    current_tags = photo.tags or ""
                    
                    file_path = _resolve_path(Path(photo.path))
                    if not file_path.exists():
                        continue
                    
                    try:
                        img_array = np.fromfile(str(file_path), np.uint8)
                        img = cv2.imdecode(img_array, cv2.IMREAD_COLOR)
                    except Exception:
                        img = cv2.imread(str(file_path))
                        
                    if img is None:
                        cursor.execute("INSERT OR IGNORE INTO processed_objects (file_id) VALUES (?)", (photo.id,))
                        processed_count += 1
                        if processed_count % 500 == 0:
                            conn.commit()
                            s.commit()
                        continue
                        
                    try:
                        img = cv2.resize(img, (224, 224))
                        img = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
                        img = img.astype(np.float32) / 255.0
                        img -= np.array([0.485, 0.456, 0.406])
                        img /= np.array([0.229, 0.224, 0.225])
                        img = img.transpose(2, 0, 1)
                        img = np.expand_dims(img, axis=0)
                        img = np.ascontiguousarray(img) # OpenCV strictly requires contiguous memory for DNN inputs!
                        
                        net.setInput(img)
                        preds = net.forward().flatten()
                        
                        # Apply Softmax to convert raw logits to proper probabilities (0.0 to 1.0)
                        exp_preds = np.exp(preds - np.max(preds))
                        probs = exp_preds / np.sum(exp_preds)

                        # Grab the top 5 highest confidence predictions
                        classIds = np.argsort(probs)[-5:][::-1]
                        new_tags = []
                        for classId in classIds:
                            confidence = probs[classId]
                            if confidence > object_threshold:
                                label = classes[classId].split(',')[0].strip().lower().replace(" ", "_")
                                new_tags.append(f"object:{label}")
                                
                        if new_tags:
                            for tag in new_tags:
                                if tag not in current_tags:
                                    current_tags = f"{current_tags} {tag}".strip()
                            photo.tags = current_tags
                    except Exception as e:
                        print(f"ERROR: Failed to classify {file_path.name}: {e}")
                        traceback.print_exc()
                        
                    cursor.execute("INSERT OR IGNORE INTO processed_objects (file_id) VALUES (?)", (photo.id,))
                    processed_count += 1
                    if processed_count % 500 == 0:
                        s.commit()
                        conn.commit()
                
                s.commit()
            conn.commit()
    except Exception as e:
        print(f"CRITICAL: Object Scanner Error: {e}")
        traceback.print_exc()
    finally:
        object_scanner_running = False
        STATE["object_scanner_stopped"] = False
        STATE["object_scanner_current_file"] = ""

@app.post("/scan-objects")
def scan_objects():
    global object_scanner_thread, object_scanner_running
    if object_scanner_running:
        raise HTTPException(status_code=400, detail="Object scanning is already in progress.")
    object_scanner_running = True
    STATE["object_scanner_stopped"] = False
    object_scanner_thread = threading.Thread(target=_scan_and_tag_objects_worker)
    object_scanner_thread.start()
    return {"message": "Object scanning started in the background."}

@app.post("/stop-scan-objects")
def stop_scan_objects():
    global object_scanner_running
    if not object_scanner_running:
        raise HTTPException(status_code=400, detail="Object scanner is not running.")
    object_scanner_running = False
    STATE["object_scanner_stopped"] = True
    return {"message": "Stopping object scanner."}

@app.post("/reset-object-scanner-progress")
def reset_object_scanner_progress():
    try:
        ai_db_path = get_ai_db_path()
        if ai_db_path.exists():
            with sqlite3.connect(ai_db_path, timeout=15) as conn:
                conn.execute("PRAGMA journal_mode=WAL;")
                conn.execute("CREATE TABLE IF NOT EXISTS processed_objects (file_id INTEGER PRIMARY KEY)")
                conn.execute("DELETE FROM processed_objects")
                conn.commit()
        return {"message": "Object scanner progress has been reset."}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Could not reset object scanner progress: {e}")

class TagUpdateRequest(BaseModel):
    file_ids: list[int]
    tags: list[str]



combined_scanner_running = False
combined_scanner_stopped = False
combined_scanner_thread = None

@app.post("/tags/add")
def add_tags(req: TagUpdateRequest):
    with SessionLocal() as s:
        # Chunk processing to avoid SQLite IN() limits and memory spikes
        for i in range(0, len(req.file_ids), 900):
            chunk = req.file_ids[i:i + 900]
            files_to_update = s.query(FileIndex.id, FileIndex.tags).filter(FileIndex.id.in_(chunk)).all()
            mappings = []
            for f_id, tags in files_to_update:
                current_tags = set((tags or "").split())
                for tag in req.tags:
                    formatted_tag = f"object:{tag}" if ":" not in tag else tag
                    current_tags.add(formatted_tag)
                new_tags_str = " ".join(sorted(current_tags))
                if new_tags_str != tags:
                    mappings.append({"id": f_id, "tags": new_tags_str})
            if mappings:
                s.bulk_update_mappings(FileIndex, mappings)
                s.commit()
    return {"status": "success"}

@app.post("/tags/remove")
def remove_tags(req: TagUpdateRequest):
    with SessionLocal() as s:
        for i in range(0, len(req.file_ids), 900):
            chunk = req.file_ids[i:i + 900]
            files_to_update = s.query(FileIndex.id, FileIndex.tags).filter(FileIndex.id.in_(chunk)).all()
            mappings = []
            for f_id, tags in files_to_update:
                if not tags:
                    continue
                current_tags = set((tags or "").split())
                for tag in req.tags:
                    formatted_tag = f"object:{tag}" if ":" not in tag else tag
                    current_tags.discard(formatted_tag)
                new_tags_str = " ".join(sorted(current_tags))
                if new_tags_str != tags:
                    mappings.append({"id": f_id, "tags": new_tags_str})
            if mappings:
                s.bulk_update_mappings(FileIndex, mappings)
                s.commit()
    return {"status": "success"}

@app.delete("/tags/objects/all")
def clear_all_object_tags():
    with SessionLocal() as s:
        # Fetch IDs first to save memory
        file_ids = [r[0] for r in s.query(FileIndex.id).filter(FileIndex.tags.like('%object:%')).all()]
        
        chunk_size = 1000
        for i in range(0, len(file_ids), chunk_size):
            chunk = file_ids[i:i + chunk_size]
            files = s.query(FileIndex.id, FileIndex.tags).filter(FileIndex.id.in_(chunk)).all()
            mappings = []
            for f_id, tags in files:
                if tags:
                    # Split tags, filter out object tags, and rejoin
                    tags_list = [t for t in re.split(r'[\s,]+', tags) if not t.startswith('object:')]
                    new_tags_str = " ".join(filter(bool, tags_list))
                    if new_tags_str != tags:
                        mappings.append({"id": f_id, "tags": new_tags_str})
            if mappings:
                s.bulk_update_mappings(FileIndex, mappings)
                s.commit()

    return {"status": "success", "message": "All object tags have been cleared."}

@app.delete("/tags/objects/{tag_name}")
def delete_object_tag_globally(tag_name: str):
    tag_to_delete = tag_name
    if not tag_to_delete.startswith("object:"):
        tag_to_delete = f"object:{tag_to_delete}"

    with SessionLocal() as s:
        file_ids = [r[0] for r in s.query(FileIndex.id).filter(FileIndex.tags.like(f'%{tag_to_delete}%')).all()]
        
        chunk_size = 1000
        for i in range(0, len(file_ids), chunk_size):
            chunk = file_ids[i:i + chunk_size]
            files = s.query(FileIndex.id, FileIndex.tags).filter(FileIndex.id.in_(chunk)).all()
            mappings = []
            for f_id, tags in files:
                if tags:
                    tags_list = [t for t in re.split(r'[\s,]+', tags) if t != tag_to_delete]
                    new_tags_str = " ".join(filter(bool, tags_list))
                    if new_tags_str != tags:
                        mappings.append({"id": f_id, "tags": new_tags_str})
            if mappings:
                s.bulk_update_mappings(FileIndex, mappings)
                s.commit()

    return {"status": "success", "deleted_tag": tag_to_delete}

@app.post("/system/free-memory")
def free_memory():
    import gc
    import sqlite3
    
    # 1. Force Python garbage collection to drop unreferenced objects
    gc.collect()
    
    # 2. Flush SQLite memory caches (FTS5 / Page Cache)
    cfg = load_config()
    db_path_str = cfg.get("database_path") or "archive.db"
    main_db_path = Path(db_path_str)
    if not main_db_path.is_absolute():
        if getattr(sys, 'frozen', False):
            main_db_path = Path(sys.executable).parent / main_db_path
        else:
            main_db_path = Path(__file__).resolve().parent.parent.parent / main_db_path

    try:
        if main_db_path.exists():
            with sqlite3.connect(main_db_path) as db:
                db.execute("PRAGMA shrink_memory")
    except Exception:
        pass
        
    ai_db_path = get_ai_db_path()
    try:
        if ai_db_path.exists():
            with sqlite3.connect(ai_db_path) as db:
                db.execute("PRAGMA shrink_memory")
    except Exception:
        pass

    if cfg.get("enable_logging"):
        import logging
        logging.info("System memory released via garbage collection and SQLite cache flush.")

    return {"status": "Memory released"}

@app.post("/system/backup")
def backup_databases(payload: dict = Body(...)):
    dest_dir = payload.get("destination")
    if not dest_dir:
        raise HTTPException(status_code=400, detail="Destination directory is required.")
        
    dest_path = Path(dest_dir)
    if not dest_path.exists() or not dest_path.is_dir():
        raise HTTPException(status_code=400, detail="Invalid destination directory.")
        
    ai_db_path = get_ai_db_path()
    cfg = load_config()
    db_path_str = cfg.get("database_path") or "archive.db"
    main_db_path = Path(db_path_str)
    if not main_db_path.is_absolute():
        if getattr(sys, 'frozen', False):
            main_db_path = Path(sys.executable).parent / main_db_path
        else:
            main_db_path = Path(__file__).resolve().parent.parent.parent / main_db_path

    if getattr(sys, 'frozen', False):
        config_path = Path(sys.executable).parent / "config.yaml"
    else:
        config_path = Path(__file__).resolve().parent.parent.parent / "config.yaml"

    try:
        if main_db_path.exists():
            with sqlite3.connect(main_db_path) as src, sqlite3.connect(dest_path / main_db_path.name) as dst:
                src.backup(dst)
        if ai_db_path.exists():
            with sqlite3.connect(ai_db_path) as src, sqlite3.connect(dest_path / ai_db_path.name) as dst:
                src.backup(dst)
        if config_path.exists():
            shutil.copy2(config_path, dest_path / config_path.name)
            
        if load_config().get("enable_logging"):
            import logging
            logging.info("Successfully backed up databases and config.")
        return {"success": True, "message": "Databases and config successfully backed up."}
    except Exception as e:
        if load_config().get("enable_logging"):
            import logging
            logging.error(f"Critical error: Backup failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Backup failed: {e}")

@app.post("/system/cleanup")
def system_cleanup():
    cfg = load_config()
    ai_db_path = get_ai_db_path()
    db_path_str = cfg.get("database_path") or "archive.db"
    main_db_path = Path(db_path_str)
    if not main_db_path.is_absolute():
        if getattr(sys, 'frozen', False):
            main_db_path = Path(sys.executable).parent / main_db_path
        else:
            main_db_path = Path(__file__).resolve().parent.parent.parent / main_db_path

    missing_ids = []
    deleted_thumbnails_count = 0
    thumb_dir = Path(cfg.get("thumbnail_path") or "thumbnails") / ".wabs_cache"
    
    with SessionLocal() as s:
        # 1. Identify files that no longer exist on disk
        for r in s.query(FileIndex.id, FileIndex.path).yield_per(1000):
            fid, path_str = r[0], r[1]
            file_path = _resolve_path(Path(path_str))
            if not file_path.exists():
                missing_ids.append(fid)
                
        # 2. Delete dead records from main archive database
        if missing_ids:
            for i in range(0, len(missing_ids), 900):
                chunk = missing_ids[i:i + 900]
                s.query(FileIndex).filter(FileIndex.id.in_(chunk)).delete(synchronize_session=False)
            s.commit()

        valid_file_ids = {str(r[0]) for r in s.query(FileIndex.id).all()}

        # 3. Clean up orphaned thumbnails
        if thumb_dir.exists():
            for f in thumb_dir.iterdir():
                if f.is_file() and f.suffix.lower() == '.jpg':
                    if f.stem not in valid_file_ids:
                        try:
                            f.unlink()
                            deleted_thumbnails_count += 1
                        except Exception:
                            pass
                            
            photos_dir = thumb_dir / "photos"
            if photos_dir.exists():
                for f in photos_dir.iterdir():
                    if f.is_file() and f.suffix.lower() == '.jpg':
                        if f.stem not in valid_file_ids:
                            try:
                                f.unlink()
                                deleted_thumbnails_count += 1
                            except Exception:
                                pass

    ai_missing_file_ids = set()
    if ai_db_path.exists():
        try:
            with sqlite3.connect(ai_db_path, timeout=15) as conn:
                cursor = conn.cursor()
                for table in ['faces', 'processed_files', 'processed_objects']:
                    try:
                        cursor.execute(f"SELECT DISTINCT file_id FROM {table}")
                        for row in cursor.fetchall():
                            if str(row[0]) not in valid_file_ids:
                                ai_missing_file_ids.add(row[0])
                    except sqlite3.OperationalError:
                        pass
        except Exception:
            pass
            
    all_missing_ids = list(set(missing_ids).union(ai_missing_file_ids))

    # 4. Clean up AI database (orphaned faces, processed files/objects, and people without faces)
    if all_missing_ids and ai_db_path.exists():
        with sqlite3.connect(ai_db_path, timeout=15) as conn:
            conn.execute("PRAGMA journal_mode=WAL;")
            cursor = conn.cursor()
            for i in range(0, len(all_missing_ids), 900):
                chunk = all_missing_ids[i:i+900]
                placeholders = ",".join("?" * len(chunk))
                
                for query in [
                    f"DELETE FROM faces WHERE file_id IN ({placeholders})",
                    f"DELETE FROM processed_files WHERE file_id IN ({placeholders})",
                    f"DELETE FROM processed_objects WHERE file_id IN ({placeholders})",
                    f"UPDATE people SET thumbnail_file_id = NULL WHERE thumbnail_file_id IN ({placeholders})"
                ]:
                    try:
                        cursor.execute(query, chunk)
                    except sqlite3.OperationalError as e:
                        if "no such table" not in str(e).lower():
                            raise
            
            try:
                # Delete people profiles that no longer have any associated faces
                cursor.execute("DELETE FROM people WHERE id NOT IN (SELECT DISTINCT person_id FROM faces)")
            except sqlite3.OperationalError as e:
                if "no such table" not in str(e).lower():
                    raise
            conn.commit()

    # 5. Clean up orphaned face thumbnails (must run after AI DB cleanup)
    if thumb_dir.exists() and ai_db_path.exists():
        faces_dir = thumb_dir / "faces"
        if faces_dir.exists():
            valid_person_ids = set()
            try:
                with sqlite3.connect(ai_db_path, timeout=15) as conn:
                    cursor = conn.cursor()
                    cursor.execute("SELECT id FROM people")
                    valid_person_ids = {str(r[0]) for r in cursor.fetchall()}
            except Exception:
                pass
            
            for f in faces_dir.iterdir():
                if f.is_file() and f.name.startswith("person_") and f.suffix.lower() == '.jpg':
                    pid_str = f.stem.replace("person_", "")
                    if pid_str not in valid_person_ids:
                        try:
                            f.unlink()
                            deleted_thumbnails_count += 1
                        except Exception:
                            pass

    # 6. Vacuum databases to reclaim space and optimize indices
    try:
        with sqlite3.connect(main_db_path, timeout=15) as conn:
            conn.isolation_level = None  # Auto-commit mode required for VACUUM
            conn.execute("VACUUM")
    except Exception as e:
        print(f"Failed to vacuum main database: {e}")
        if cfg.get("enable_logging"):
            import logging
            logging.error(f"Failed to vacuum main database: {e}", exc_info=True)

    if ai_db_path.exists():
        try:
            with sqlite3.connect(ai_db_path, timeout=15) as conn:
                conn.isolation_level = None
                conn.execute("VACUUM")
        except Exception as e:
            print(f"Failed to vacuum AI database: {e}")
            if cfg.get("enable_logging"):
                import logging
                logging.error(f"Failed to vacuum AI database: {e}", exc_info=True)

    if cfg.get("enable_logging"):
        import logging
        logging.info(f"Database cleanup finished: {len(missing_ids)} dead files removed, {deleted_thumbnails_count} orphaned thumbnails deleted, empty profiles cleared, and databases vacuumed.")

    return {"status": "success", "removed_files": len(missing_ids), "removed_thumbnails": deleted_thumbnails_count, "message": "Cleanup and optimization complete."}

@app.post("/system/purge-unknowns")
def purge_unknowns(payload: dict = Body(...)):
    # Purge Unknown profiles with fewer faces than the threshold (default: 3)
    threshold = int(payload.get("threshold", 3))
    cfg = load_config()
    ai_db_path = get_ai_db_path()
    
    if not ai_db_path.exists():
        raise HTTPException(status_code=404, detail="AI Database not found")
        
    purged_count = 0
    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        
        # 1. Find all Unknown People with faces less than the threshold
        cursor.execute("""
            SELECT p.id 
            FROM people p 
            LEFT JOIN faces f ON p.id = f.person_id 
            WHERE p.name LIKE 'Unknown Person%' 
            GROUP BY p.id 
            HAVING COUNT(f.id) < ?
        """, (threshold,))
        
        ids_to_delete = [r[0] for r in cursor.fetchall()]
        
        if ids_to_delete:
            for i in range(0, len(ids_to_delete), 900):
                chunk = ids_to_delete[i:i+900]
                placeholders = ",".join("?" * len(chunk))
                
                # 2. Delete the heavy embeddings, face records, and person profile
                cursor.execute(f"DELETE FROM faces WHERE person_id IN ({placeholders})", chunk)
                cursor.execute(f"DELETE FROM people WHERE id IN ({placeholders})", chunk)
                
                # Note: We intentionally DO NOT delete from 'processed_files'.
                # This guarantees the AI scanner will permanently skip these photos in the future.
                for pid in chunk:
                    EXEMPLAR_CACHE.pop(pid)
                    
            purged_count = len(ids_to_delete)
            conn.commit()
            
    # Trigger native system cleanup to vacuum the DB and delete orphaned thumbnails
    system_cleanup()
    
    if cfg.get("enable_logging"):
        import logging
        logging.info(f"Purged {purged_count} small unknown profiles to reclaim space.")
        
    return {"status": "success", "purged_profiles": purged_count}

@app.get("/system/export-people")
def export_people():
    if load_config().get("enable_logging"):
        import logging
        logging.info("Exporting people data.")
    ai_db_path = get_ai_db_path()
    if not ai_db_path.exists():
        return []
         
    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT id, name, thumbnail_file_id FROM people WHERE name NOT LIKE 'Unknown Person%'")
            people_rows = cursor.fetchall()
        except sqlite3.OperationalError as e:
            if "no such table" not in str(e).lower():
                raise HTTPException(status_code=500, detail=f"Database locked or unavailable: {e}")
            people_rows = []
        
        export_data = []
        with SessionLocal() as s:
            for pid, name, thumb_id in people_rows:
                thumb_path = None
                if thumb_id:
                    thumb_file = s.get(FileIndex, thumb_id)
                    if thumb_file:
                        thumb_path = thumb_file.path
                        
                cursor.execute("SELECT file_id, embedding_json FROM faces WHERE person_id = ?", (pid,))
                face_rows = cursor.fetchall()
                
                faces = []
                for fid, emb_json in face_rows:
                    f_item = s.get(FileIndex, fid)
                    if f_item:
                        faces.append({
                            "path": f_item.path,
                            "embedding": emb_json
                        })
                        
                if faces:
                    export_data.append({
                        "name": name,
                        "thumbnail_path": thumb_path,
                        "faces": faces
                    })
    return export_data

@app.post("/system/import-people")
def import_people(payload: list = Body(...)):
    ai_db_path = get_ai_db_path()
    with sqlite3.connect(ai_db_path, timeout=15) as conn:
        conn.execute("PRAGMA journal_mode=WAL;")
        cursor = conn.cursor()
        cursor.execute('''CREATE TABLE IF NOT EXISTS people (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                            name TEXT DEFAULT 'Unknown Person',
                            thumbnail_file_id INTEGER
                      )''')
        cursor.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_people_name ON people(name)")
        cursor.execute('''CREATE TABLE IF NOT EXISTS faces (
                            id INTEGER PRIMARY KEY AUTOINCREMENT,
                            person_id INTEGER,
                            file_id INTEGER,
                            embedding_json TEXT,
                            FOREIGN KEY(person_id) REFERENCES people(id)
                        )''')
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_faces_person_file ON faces(person_id, file_id)")
        cursor.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_faces_unique ON faces(person_id, file_id, embedding_json)")
        cursor.execute('''CREATE TABLE IF NOT EXISTS processed_files (file_id INTEGER PRIMARY KEY)''')
        
        imported_people = 0
        imported_faces = 0
        with SessionLocal() as s:
            paths_in_payload = set()
            for p_data in payload:
                if t := p_data.get("thumbnail_path"):
                    paths_in_payload.add(t)
                for face in p_data.get("faces", []):
                    if f := face.get("path"):
                        paths_in_payload.add(f)
            
            path_to_id = {}
            if paths_in_payload:
                paths_list = list(paths_in_payload)
                for i in range(0, len(paths_list), 900):
                    chunk = paths_list[i:i+900]
                    for r in s.query(FileIndex.path, FileIndex.id).filter(FileIndex.path.in_(chunk)).all():
                        path_to_id[r[0]] = r[1]
            
            fallback_map = None
            def get_fid(path_str):
                if not path_str: return None
                if path_str in path_to_id: return path_to_id[path_str]
                
                nonlocal fallback_map
                if fallback_map is None:
                    fallback_map = {}
                    for r in s.query(FileIndex.path, FileIndex.id).yield_per(10000):
                        p, fid = r[0], r[1]
                        try:
                            p_clean = p.replace('\\', '/')
                            parts = p_clean.strip('/').split('/')
                            key = f"{parts[-2]}/{parts[-1]}" if len(parts) >= 2 else parts[-1]
                            if key not in fallback_map: fallback_map[key] = []
                            fallback_map[key].append(fid)
                        except Exception:
                            pass
                
                try:
                    p_clean = path_str.replace('\\', '/')
                    parts = p_clean.strip('/').split('/')
                    key = f"{parts[-2]}/{parts[-1]}" if len(parts) >= 2 else parts[-1]
                    matches = fallback_map.get(key, [])
                    if len(matches) == 1: 
                        path_to_id[path_str] = matches[0]
                        return matches[0]
                except Exception:
                    pass
                return None

            tags_to_update = {}

            for p_data in payload:
                name = p_data.get("name")
                if not name: continue
                cursor.execute("INSERT OR IGNORE INTO people (name) VALUES (?)", (name,))
                if cursor.rowcount > 0:
                    imported_people += 1
                cursor.execute("SELECT id FROM people WHERE name = ?", (name,))
                pid = cursor.fetchone()[0]
                
                thumb_path = p_data.get("thumbnail_path")
                thumb_fid = get_fid(thumb_path)
                for face in p_data.get("faces", []):
                    f_path = face.get("path")
                    emb = face.get("embedding")
                    fid = get_fid(f_path)
                    if fid:
                        cursor.execute("INSERT OR IGNORE INTO faces (person_id, file_id, embedding_json) VALUES (?, ?, ?)", (pid, fid, emb))
                        if cursor.rowcount > 0:
                            cursor.execute("INSERT OR IGNORE INTO processed_files (file_id) VALUES (?)", (fid,))
                            imported_faces += 1
                            if fid not in tags_to_update:
                                tags_to_update[fid] = set()
                            tags_to_update[fid].add(f"person:{name}")
                
                if thumb_fid:
                    cursor.execute("UPDATE people SET thumbnail_file_id = ? WHERE id = ?", (thumb_fid, pid))

            if tags_to_update:
                fid_list = list(tags_to_update.keys())
                for i in range(0, len(fid_list), 900):
                    chunk = fid_list[i:i+900]
                    files_to_update = s.query(FileIndex.id, FileIndex.tags).filter(FileIndex.id.in_(chunk)).all()
                    mappings = []
                    for f_id, tags in files_to_update:
                        current_tags = set((tags or "").split())
                        new_tags = tags_to_update[f_id]
                        new_tags_str = " ".join(sorted(current_tags.union(new_tags)))
                        if new_tags_str != tags:
                            mappings.append({"id": f_id, "tags": new_tags_str})
                    if mappings:
                        s.bulk_update_mappings(FileIndex, mappings)

            s.commit()
        conn.commit()
    return {"success": True, "imported_people": imported_people, "imported_faces": imported_faces}

@app.get("/tags/objects")
def get_object_tags():
    with SessionLocal() as s:
        unique_tags = set()
        # Use yield_per to stream results instead of loading all strings into memory
        for r in s.query(FileIndex.tags).filter(FileIndex.tags.like('%object:%')).yield_per(1000):
            if r[0]:
                for tag in r[0].split():
                    if tag.startswith('object:'):
                        unique_tags.add(tag)
        return sorted(list(unique_tags))

@app.get("/system/export-tags")
def export_tags():
    if load_config().get("enable_logging"):
        import logging
        logging.info("Exporting tags data.")
    with SessionLocal() as s:
        # Stream all files that have any tags
        files = s.query(FileIndex.path, FileIndex.tags).filter(FileIndex.tags != None, FileIndex.tags != '').yield_per(1000)
        export_data = [{"path": path, "tags": tags} for path, tags in files]
        return export_data

@app.post("/system/import-tags")
def import_tags(payload: list = Body(...)):
    imported_count = 0
    with SessionLocal() as s:
        paths_in_payload = set()
        for item in payload:
            if p := item.get("path"):
                paths_in_payload.add(p)
        
        path_to_id = {}
        if paths_in_payload:
            paths_list = list(paths_in_payload)
            for i in range(0, len(paths_list), 900):
                chunk = paths_list[i:i+900]
                for r in s.query(FileIndex.path, FileIndex.id).filter(FileIndex.path.in_(chunk)).all():
                    path_to_id[r[0]] = r[1]

        fallback_map = None
        def get_fid(path_str):
            if not path_str: return None
            if path_str in path_to_id: return path_to_id[path_str]
            
            nonlocal fallback_map
            if fallback_map is None:
                fallback_map = {}
                for r in s.query(FileIndex.path, FileIndex.id).yield_per(10000):
                    p, fid = r[0], r[1]
                    try:
                        p_clean = p.replace('\\', '/')
                        parts = p_clean.strip('/').split('/')
                        key = f"{parts[-2]}/{parts[-1]}" if len(parts) >= 2 else parts[-1]
                        if key not in fallback_map: fallback_map[key] = []
                        fallback_map[key].append(fid)
                    except Exception:
                        pass
            
            try:
                p_clean = path_str.replace('\\', '/')
                parts = p_clean.strip('/').split('/')
                key = f"{parts[-2]}/{parts[-1]}" if len(parts) >= 2 else parts[-1]
                matches = fallback_map.get(key, [])
                if len(matches) == 1: 
                    path_to_id[path_str] = matches[0]
                    return matches[0]
            except Exception:
                pass
            return None
            
        chunk_size = 900
        for i in range(0, len(payload), chunk_size):
            chunk = payload[i:i + chunk_size]
            chunk_updates = {}
            
            for item in chunk:
                path = item.get("path")
                new_tags = item.get("tags")
                if not path or not new_tags: 
                    continue
                
                file_id = get_fid(path)
                if file_id:
                    chunk_updates[file_id] = new_tags
                    
            if not chunk_updates:
                continue
                
            files_to_update = s.query(FileIndex.id, FileIndex.tags).filter(FileIndex.id.in_(chunk_updates.keys())).all()
            mappings = []
            for f_id, tags in files_to_update:
                current_tags = set((tags or "").split())
                imported_tags = set(chunk_updates[f_id].split())
                new_tags_str = " ".join(sorted(current_tags.union(imported_tags)))
                if new_tags_str != tags:
                    mappings.append({"id": f_id, "tags": new_tags_str})
                    imported_count += 1
            
            if mappings:
                s.bulk_update_mappings(FileIndex, mappings)
                s.commit()
    return {"success": True, "imported_files": imported_count}

# --- Serve React Frontend (Production) ---
if hasattr(sys, '_MEIPASS'):
    # Running as a packaged PyInstaller executable
    frontend_dist = Path(sys._MEIPASS) / "frontend" / "dist"
else:
    # Running in normal development mode
    frontend_dist = Path(__file__).parent.parent.parent / "frontend" / "dist"

if frontend_dist.exists():
    app.mount("/assets", StaticFiles(directory=str(frontend_dist / "assets")), name="assets")
    @app.get("/{full_path:path}")
    def serve_frontend(full_path: str):
        path = frontend_dist / full_path
        if path.exists() and path.is_file():
            return FileResponse(str(path))
        return FileResponse(str(frontend_dist / "index.html"))